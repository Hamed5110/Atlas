"""Live end-to-end smoke: every screen API + import/export/print surfaces.

Uses in-process TestClient (SQLite) so no .env credentials are required.
Reports a per-module pass/fail matrix to stdout.
"""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from airfare_management.api.main import create_app
from airfare_management.config import Settings
from tests.helpers import create_employee, ensure_company


def _xlsx_bytes(headers: list[str], rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> int:
    settings = Settings(
        environment="test",
        database_url="sqlite+pysqlite:///:memory:",
        jwt_secret="t" * 32,
        bootstrap_admin_password="StrongPassword!2026",
        attachment_root="./var/test-smoke-attachments",
    )
    results: list[tuple[str, str, bool, str]] = []

    def check(module: str, name: str, ok: bool, detail: str = "") -> None:
        safe = detail.encode("ascii", "replace").decode("ascii")
        results.append((module, name, ok, safe))
        mark = "PASS" if ok else "FAIL"
        print(f"[{mark}] {module} :: {name}" + (f" -- {safe}" if safe else ""))

    with TestClient(create_app(settings)) as client:
        login = client.post(
            "/v1/auth/login",
            json={"username": "admin", "password": "StrongPassword!2026"},
        )
        check("auth", "login", login.status_code == 200, str(login.status_code))
        if login.status_code != 200:
            return 1
        token = login.json()["access_token"]
        h = {"Authorization": f"Bearer {token}"}

        me = client.get("/v1/auth/me", headers=h)
        check("auth", "me", me.status_code == 200)

        screen_gets = [
            ("dashboard", "/v1/dashboard"),
            ("employees", "/v1/employees"),
            ("opening-balances", "/v1/opening-balances"),
            ("loans", "/v1/loans"),
            ("ess", "/v1/ess/requests"),
            ("ess-dashboard", "/v1/ess/dashboard"),
            ("rates", "/v1/entitlement-rates"),
            ("tickets", "/v1/tickets"),
            ("lookups", "/v1/lookup-types"),
            ("users", "/v1/users"),
            ("audit", "/v1/audit"),
            ("settings", "/v1/settings"),
            ("preferences", "/v1/preferences/effective"),
            ("companies", "/v1/companies"),
            ("documents", "/v1/documents"),
            ("document-templates", "/v1/documents/templates"),
            ("report-templates", "/v1/report-templates"),
            ("ai-learning", "/v1/ai/support/learning"),
            ("ai-forecast", "/v1/ai/forecasts/budget"),
            ("ai-schema", "/v1/ai/agent/schema"),
            ("backups", "/v1/admin/backups"),
            ("templates-list", "/v1/templates"),
            ("crystal-status", "/v1/crystal/status"),
        ]
        for module, path in screen_gets:
            r = client.get(path, headers=h)
            check(module, f"GET {path}", r.status_code in {200, 204}, f"status={r.status_code}")

        for name in ("employees", "opening-balances"):
            r = client.get(f"/v1/templates/{name}.xlsx", headers=h)
            ok = r.status_code == 200 and r.content[:2] == b"PK"
            check(
                "templates",
                f"download {name}.xlsx",
                ok,
                f"status={r.status_code} bytes={len(r.content)}",
            )

        r = client.get("/v1/employees/export.xlsx", headers=h)
        check(
            "employees",
            "export.xlsx",
            r.status_code == 200 and r.content[:2] == b"PK",
            f"status={r.status_code}",
        )

        r = client.get("/v1/opening-balances/export.xlsx", headers=h)
        check(
            "opening-balances",
            "export.xlsx",
            r.status_code == 200 and r.content[:2] == b"PK",
            f"status={r.status_code}",
        )

        company_id = ensure_company(client, h)
        employee = create_employee(client, h, company_id=company_id, code="SMOKE001")
        emp_id = employee["id"]
        check("employees", "create", bool(emp_id), str(emp_id))

        tmpl = client.get("/v1/templates/employees.xlsx", headers=h)
        if tmpl.status_code == 200:
            wb = load_workbook(io.BytesIO(tmpl.content))
            ws = wb.active
            col_headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            row_map = {hname: "" for hname in col_headers if hname}
            # Force required template keys when present (exact column names)
            for hname in list(row_map):
                low = str(hname).lower().strip()
                if low == "company_id":
                    row_map[hname] = company_id
                elif low == "join_date":
                    row_map[hname] = "2024-01-15"
                elif low == "code":
                    row_map[hname] = "SMOKEIMP01"
                elif low == "full_name":
                    row_map[hname] = "Import Smoke"
                elif low == "passport_no":
                    row_map[hname] = "PX999"
                elif low.endswith("_expiry") or low.endswith("_date"):
                    # leave blank dates empty (None) — do not fill with passport text
                    row_map[hname] = None if row_map[hname] in ("", "PX999") else row_map[hname]
            payload = _xlsx_bytes(list(row_map.keys()), [list(row_map.values())])
            preview = client.post(
                "/v1/employees/import/preview",
                headers=h,
                files={
                    "file": (
                        "employees.xlsx",
                        payload,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            check(
                "employees",
                "import/preview",
                preview.status_code == 200,
                f"status={preview.status_code}",
            )
            if preview.status_code == 200:
                body = preview.json()
                rows = body.get("rows", body) if isinstance(body, dict) else body
                ready = [r for r in rows if isinstance(r, dict) and r.get("severity") == "READY"]
                if ready:
                    for rr in ready:
                        rr["selected"] = True
                    commit = client.post(
                        "/v1/employees/import/commit",
                        headers=h,
                        json={"rows": ready},
                    )
                    check(
                        "employees",
                        "import/commit",
                        commit.status_code == 200,
                        f"status={commit.status_code} {commit.text[:100]}",
                    )
                else:
                    check(
                        "employees",
                        "import/commit",
                        False,
                        f"no READY rows: {json.dumps(rows)[:240]}",
                    )

        ob_tmpl = client.get("/v1/templates/opening-balances.xlsx", headers=h)
        if ob_tmpl.status_code == 200:
            wb = load_workbook(io.BytesIO(ob_tmpl.content))
            ws = wb.active
            col_headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            row_map = {hname: "" for hname in col_headers if hname}
            for key, val in {
                "employee code": "SMOKE001",
                "code": "SMOKE001",
                "days": 5,
                "amount": 100,
                "balance days": 5,
                "opening": 5,
            }.items():
                for hname in list(row_map):
                    if key in str(hname).lower():
                        row_map[hname] = val
            payload = _xlsx_bytes(list(row_map.keys()), [list(row_map.values())])
            preview = client.post(
                "/v1/opening-balances/import/preview",
                headers=h,
                files={
                    "file": (
                        "ob.xlsx",
                        payload,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            check(
                "opening-balances",
                "import/preview",
                preview.status_code == 200,
                f"status={preview.status_code}",
            )

        for dataset in (
            "employee-master",
            "opening-balances",
            "ticket-register",
            "loan-outstanding",
            "excess-recovery",
            "entitlements",
            "loan-statement",
            "liability-projections",
        ):
            for fmt in ("xlsx", "pdf"):
                r = client.get(f"/v1/reports/export/{dataset}.{fmt}", headers=h)
                ok = r.status_code == 200 and len(r.content) > 20
                check(
                    "reports",
                    f"export {dataset}.{fmt}",
                    ok,
                    f"status={r.status_code} bytes={len(r.content)}",
                )

        # Seed entitlement rate so allocation print can render
        rate = client.post(
            "/v1/entitlement-rates",
            headers=h,
            json={
                "scope_type": "global",
                "scope_id": "",
                "amount": "1.0000",
                "effective_from": "2020-01-01",
                "effective_to": None,
                "cap_amount": "500",
            },
        )
        check("rates", "seed for print", rate.status_code in {200, 201}, f"status={rate.status_code} {rate.text[:100]}")

        preview = client.post(
            "/v1/allocations/preview",
            headers=h,
            json={
                "employee_id": emp_id,
                "requested_ticket_amount": 100,
                "as_of_date": "2026-09-01",
            },
        )
        check(
            "allocation",
            "preview",
            preview.status_code == 200,
            f"status={preview.status_code} {preview.text[:140]}",
        )

        for path in ("/v1/allocations/print", "/v1/allocations/print.pdf"):
            r = client.post(
                path,
                headers=h,
                json={
                    "employee_id": emp_id,
                    "requested_ticket_amount": 100,
                    "as_of_date": "2026-09-01",
                },
            )
            ok = r.status_code == 200 and (
                r.content[:4] == b"%PDF"
                or "pdf" in (r.headers.get("content-type") or "").lower()
            )
            check(
                "allocation",
                path.split("/")[-1],
                ok,
                f"status={r.status_code} ctype={r.headers.get('content-type')} bytes={len(r.content)}",
            )

        templates = client.get("/v1/documents/templates", headers=h)
        check("documents", "templates", templates.status_code == 200)
        if templates.status_code == 200:
            try:
                raw = templates.json()
                tpls = raw.get("templates", []) if isinstance(raw, dict) else raw
                if not isinstance(tpls, list):
                    tpls = []
                offer = next(
                    (
                        t
                        for t in tpls
                        if isinstance(t, dict)
                        and (
                            "offer" in str(t.get("kind", "")).lower()
                            or "offer" in str(t.get("key", "")).lower()
                        )
                    ),
                    tpls[0] if tpls else None,
                )
                if not offer:
                    check("documents", "issue", False, "no templates")
                else:
                    issued = client.post(
                        "/v1/documents",
                        headers=h,
                        json={
                            "kind": offer.get("kind", "offer_letter"),
                            "template_key": offer.get("key"),
                            "employee_id": emp_id,
                            "params": {
                                "nature_of_employment": "Permanent",
                                "joining_date": "2026-09-01",
                                "document_date": "2026-09-01",
                                "probation_months": "3",
                                "basic": "500",
                                "annual_leave_days": "30",
                                "offer_valid_until": "2026-09-15",
                            },
                        },
                    )
                    check(
                        "documents",
                        "issue",
                        issued.status_code in {200, 201},
                        f"status={issued.status_code} {issued.text[:140]}",
                    )
                    if issued.status_code in {200, 201}:
                        doc_id = issued.json().get("id")
                        pdf = client.get(f"/v1/documents/{doc_id}/pdf", headers=h)
                        check(
                            "documents",
                            "pdf print layout",
                            pdf.status_code == 200 and pdf.content[:4] == b"%PDF",
                            f"status={pdf.status_code} bytes={len(pdf.content)}",
                        )
            except Exception as exc:  # noqa: BLE001
                check("documents", "issue", False, f"exception: {exc}")

        listed = client.get("/v1/report-templates", headers=h)
        if listed.status_code == 200 and listed.json():
            tid = listed.json()[0]["id"]
            for fmt in ("xlsx", "pdf"):
                r = client.get(f"/v1/report-templates/{tid}/run?format={fmt}", headers=h)
                check(
                    "reports",
                    f"template-run.{fmt}",
                    r.status_code == 200 and len(r.content) > 20,
                    f"status={r.status_code} bytes={len(r.content)}",
                )

        for path, method, body in (
            ("/v1/ai/support/diagnose", "GET", None),
            ("/v1/ai/anomalies", "POST", {}),
            ("/v1/ai/agent/chat", "POST", {"message": "Run diagnostics"}),
            ("/v1/ai/saa/baseline", "POST", {}),
        ):
            if method == "GET":
                r = client.get(path, headers=h)
            else:
                r = client.post(path, headers=h, json=body)
            check("ai", path, r.status_code == 200, f"status={r.status_code}")

        web_root = Path(r"C:\HCM Airfare\src\airfare_management\interface\web_dist_next")
        screens = [
            "dashboard",
            "allocation",
            "employees",
            "opening-balances",
            "loans",
            "ess",
            "rates",
            "reports",
            "ai-insights",
            "offer-letters",
            "contracts",
            "settings",
            "lookups",
            "users",
            "backups",
            "audit",
            "login",
        ]
        for screen in screens:
            html = web_root / screen / "index.html"
            if screen == "login":
                src = Path(r"C:\Airfare_Allowance\atlas-next\app\login\page.tsx")
            else:
                src = Path(r"C:\Airfare_Allowance\atlas-next\app") / "(app)" / screen / "page.tsx"
            deployed = html.is_file()
            sourced = src.is_file()
            ok = deployed if screen == "backups" else (deployed and sourced)
            check(
                "frontend",
                screen,
                ok,
                f"deployed={deployed} source={sourced}",
            )

    passed = sum(1 for *_, ok, _ in results if ok)
    failed = sum(1 for *_, ok, _ in results if not ok)
    print("\n=== SUMMARY ===")
    print(f"passed={passed} failed={failed} total={len(results)}")
    fails = [r for r in results if not r[2]]
    if fails:
        print("FAILURES:")
        for module, name, _, detail in fails:
            print(f"  - {module} :: {name} :: {detail}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())

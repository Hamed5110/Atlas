from __future__ import annotations

import csv
import hashlib
import io
import json
import secrets
from datetime import date, datetime, timezone
from datetime import timedelta
from decimal import Decimal
from enum import StrEnum
from typing import Annotated, Any

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, status
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.database import get_db


router = APIRouter(prefix="/api/v1", tags=["Module 2-11 Blueprints"])
SessionDep = Annotated[Session, Depends(get_db)]


class ImportModule(StrEnum):
    EMPLOYEES = "employees"


class ImportPreviewResponse(BaseModel):
    total: int
    valid: int
    failed: int
    errors: list[dict[str, Any]]
    rows: list[dict[str, Any]]


class ImportFieldStatus(BaseModel):
    column: str
    value: str | None
    status: str = Field(pattern="^(VALID|WARNING|ERROR)$")
    reason: str | None = None


class ImportPreviewRecord(BaseModel):
    row: int
    status: str = Field(pattern="^(VALID|WARNING|ERROR)$")
    fields: list[ImportFieldStatus]
    data: dict[str, Any]


class ImportVerifyPreviewResponse(BaseModel):
    previewToken: str
    expiresAtUtc: datetime
    totalRows: int
    validRowsCount: int
    errorRowsCount: int
    duplicateRowsCount: int
    previewData: list[ImportPreviewRecord]
    errorDetails: list[dict[str, Any]]


class ImportCommitRequest(BaseModel):
    previewToken: str = Field(min_length=32, max_length=96)


class ImportCommitResponse(BaseModel):
    previewToken: str
    insertedRows: int
    skippedRows: int
    status: str


class AirfareEntitlementResponse(BaseModel):
    employeeId: int
    employeeCode: str
    fullName: str
    joiningDate: date
    targetDate: date
    elapsedServiceDays: int
    monthlyRate: Decimal
    accruedBalance: Decimal
    seedBalance: Decimal
    claimedBalance: Decimal
    availableBalance: Decimal
    model: str


class AirfareClaimCreate(BaseModel):
    employee_id: int = Field(ge=1)
    claim_date: date
    sector_code: str = Field(min_length=1, max_length=40)
    claim_type: str = Field(pattern="^(Ticket|CashEncashment|DependentTicket)$")
    claim_amount: Decimal = Field(ge=0, max_digits=18, decimal_places=3)
    accrued_balance: Decimal = Field(ge=0, max_digits=18, decimal_places=3)
    notes: str | None = Field(default=None, max_length=400)

    @model_validator(mode="after")
    def claim_cannot_exceed_balance(self) -> "AirfareClaimCreate":
        if self.claim_amount > self.accrued_balance:
            raise ValueError("claim_amount cannot exceed accrued_balance.")
        return self


class LoanRequestCreate(BaseModel):
    employee_id: int = Field(ge=1)
    principal_amount: Decimal = Field(gt=0, max_digits=18, decimal_places=3)
    terms_months: int = Field(ge=1, le=120)
    annual_interest_rate: Decimal = Field(default=Decimal("0.000"), ge=0, le=100, max_digits=9, decimal_places=3)
    disbursement_date: date
    notes: str | None = Field(default=None, max_length=400)


class SeedEvidenceCreate(BaseModel):
    employee_id: int = Field(ge=1)
    seed_type: str = Field(pattern="^(LeaveBalance|GratuityAccrual|SalaryAdjustment|AirfareEntitlement)$")
    effective_date: date
    amount: Decimal = Field(max_digits=18, decimal_places=3)
    source_reference: str = Field(min_length=1, max_length=120)
    approved_by: str = Field(min_length=1, max_length=120)


class ReportRequest(BaseModel):
    report_code: str = Field(pattern="^(PayrollSummary|DepartmentCosting|DocumentExpiry|LoanBalances|AirfareUtilization)$")
    date_from: date | None = None
    date_to: date | None = None
    department_id: int | None = Field(default=None, ge=1)
    employee_id: int | None = Field(default=None, ge=1)
    export_format: str = Field(default="json", pattern="^(json|xlsx|pdf)$")


class SystemSettingUpsert(BaseModel):
    setting_key: str = Field(min_length=1, max_length=120)
    setting_value: str = Field(max_length=4000)
    value_type: str = Field(default="string", pattern="^(string|number|boolean|json|secret)$")
    is_secret: bool = False


class SelfServiceRequestCreate(BaseModel):
    employee_id: int = Field(ge=1)
    request_type: str = Field(pattern="^(PayslipView|LoanRequest|AirfareRequest|LeaveStatus|DocumentStatus|ProfileUpdate)$")
    requested_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    payload: dict[str, Any] = Field(default_factory=dict)


class AirportResponse(BaseModel):
    iata_code: str
    airport_name: str
    city_name: str
    country_code: str
    sector_code: str | None
    score: int


@router.post("/imports/{module_code}/preview", response_model=ImportPreviewResponse)
async def preview_import(module_code: ImportModule, file: UploadFile) -> ImportPreviewResponse:
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Import file exceeds 10 MB.")
    rows = parse_csv_bytes(content)
    errors = []
    valid_rows = []
    for index, row in enumerate(rows, start=1):
        row_errors = validate_employee_import_row(row)
        if row_errors:
            errors.append({"row": index, "errors": row_errors, "data": row})
        else:
            valid_rows.append(row)
    return ImportPreviewResponse(total=len(rows), valid=len(valid_rows), failed=len(errors), errors=errors, rows=valid_rows)


@router.post("/import/verify-preview", response_model=ImportVerifyPreviewResponse)
async def verify_import_preview(file: UploadFile, session: SessionDep, module_code: ImportModule = ImportModule.EMPLOYEES) -> ImportVerifyPreviewResponse:
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Import file exceeds 10 MB.")

    rows = parse_import_file(file.filename or "import.csv", content)
    existing_codes = {
        str(value).upper()
        for value in session.execute(text("SELECT EmployeeCode FROM core.Employees WHERE IsDeleted = 0")).scalars().all()
    }
    seen_codes: set[str] = set()
    preview_records: list[ImportPreviewRecord] = []
    error_details: list[dict[str, Any]] = []
    valid_rows: list[dict[str, Any]] = []
    duplicate_rows = 0

    for row_number, row in enumerate(rows, start=1):
        field_statuses, row_errors, is_duplicate = validate_employee_preview_row(row, row_number, existing_codes, seen_codes)
        if is_duplicate:
            duplicate_rows += 1
        row_status = "ERROR" if row_errors else ("WARNING" if any(field.status == "WARNING" for field in field_statuses) else "VALID")
        for detail in row_errors:
            error_details.append(detail)
        if not row_errors:
            valid_rows.append(normalize_employee_import_row(row))
        preview_records.append(ImportPreviewRecord(row=row_number, status=row_status, fields=field_statuses, data=row))

    token = secrets.token_urlsafe(48)
    expires_at = datetime.now(timezone.utc) + timedelta(hours=2)
    payload = {
        "moduleCode": module_code.value,
        "fileName": file.filename or "import.csv",
        "validRows": valid_rows,
        "errors": error_details,
        "contentSha256": hashlib.sha256(content).hexdigest(),
    }
    session.execute(
        text(
            """
            INSERT INTO core.ImportPreviewSessions
                (PreviewToken, ModuleCode, FileName, PayloadJSON, TotalRows, ValidRowsCount, ErrorRowsCount, DuplicateRowsCount, ExpiresAtUtc)
            VALUES
                (:token, :module_code, :file_name, :payload_json, :total, :valid, :errors, :duplicates, :expires_at)
            """
        ),
        {
            "token": token,
            "module_code": module_code.value,
            "file_name": file.filename or "import.csv",
            "payload_json": json.dumps(payload, separators=(",", ":"), default=str),
            "total": len(rows),
            "valid": len(valid_rows),
            "errors": len(error_details),
            "duplicates": duplicate_rows,
            "expires_at": expires_at.replace(tzinfo=None),
        },
    )
    session.commit()

    return ImportVerifyPreviewResponse(
        previewToken=token,
        expiresAtUtc=expires_at,
        totalRows=len(rows),
        validRowsCount=len(valid_rows),
        errorRowsCount=len(error_details),
        duplicateRowsCount=duplicate_rows,
        previewData=preview_records[:50],
        errorDetails=error_details,
    )


@router.post("/import/commit", response_model=ImportCommitResponse)
def commit_import(payload: ImportCommitRequest, session: SessionDep) -> ImportCommitResponse:
    record = session.execute(
        text(
            """
            SELECT PreviewToken, PayloadJSON, IsCommitted, ExpiresAtUtc
            FROM core.ImportPreviewSessions WITH (UPDLOCK, ROWLOCK)
            WHERE PreviewToken = :token
            """
        ),
        {"token": payload.previewToken},
    ).mappings().first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Preview token not found.")
    if bool(record["IsCommitted"]):
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Preview token was already committed.")
    if record["ExpiresAtUtc"] < datetime.utcnow():
        raise HTTPException(status_code=status.HTTP_410_GONE, detail="Preview token expired. Re-run verification.")

    stored_payload = json.loads(record["PayloadJSON"])
    valid_rows = stored_payload.get("validRows", [])
    inserted = 0
    skipped = 0
    batch_id = session.execute(
        text(
            """
            INSERT INTO core.ImportBatches (ModuleCode, FileName, TotalRows, BatchStatus, CreatedBy)
            OUTPUT INSERTED.ImportBatchID
            VALUES (:module_code, :file_name, :total_rows, N'Processing', N'import-preview')
            """
        ),
        {"module_code": stored_payload.get("moduleCode", "employees"), "file_name": stored_payload.get("fileName", "import.csv"), "total_rows": len(valid_rows)},
    ).scalar_one()

    try:
        for row in valid_rows:
            exists = session.execute(text("SELECT 1 FROM core.Employees WHERE EmployeeCode = :code AND IsDeleted = 0"), {"code": row["EmployeeCode"]}).first()
            if exists:
                skipped += 1
                continue
            session.execute(EMPLOYEE_IMPORT_INSERT_SQL, row)
            inserted += 1
        session.execute(
            text(
                """
                UPDATE core.ImportPreviewSessions SET IsCommitted = 1, CommittedAtUtc = SYSUTCDATETIME() WHERE PreviewToken = :token;
                UPDATE core.ImportBatches SET InsertedRows = :inserted, FailedRows = :skipped, BatchStatus = N'Completed' WHERE ImportBatchID = :batch_id;
                """
            ),
            {"token": payload.previewToken, "inserted": inserted, "skipped": skipped, "batch_id": batch_id},
        )
        session.commit()
    except Exception:
        session.rollback()
        raise

    return ImportCommitResponse(previewToken=payload.previewToken, insertedRows=inserted, skippedRows=skipped, status="Committed")


@router.get("/airfare/entitlement/{emp_id}", response_model=AirfareEntitlementResponse)
def get_airfare_entitlement(emp_id: int, session: SessionDep, target_date: Annotated[date, Query()]) -> AirfareEntitlementResponse:
    row = session.execute(
        text(
            """
            WITH RateSetting AS (
                SELECT TOP (1) TRY_CONVERT(DECIMAL(18,3), SettingValue) AS MonthlyRate
                FROM core.SystemSettings
                WHERE SettingKey = N'airfare.monthly_rate_bhd'
            ),
            ApprovedClaims AS (
                SELECT EmployeeID, SUM(ClaimAmount) AS ClaimedBalance
                FROM core.AirfareClaims
                WHERE ApprovalStatus = N'Approved' AND ClaimDate <= :target_date
                GROUP BY EmployeeID
            ),
            ApprovedSeeds AS (
                SELECT EmployeeID, SUM(Amount) AS SeedBalance
                FROM core.SeedEvidence
                WHERE SeedType = N'AirfareEntitlement' AND ApprovalStatus = N'Approved' AND EffectiveDate <= :target_date
                GROUP BY EmployeeID
            )
            SELECT
                e.EmployeeID,
                e.EmployeeCode,
                e.FullName,
                e.JoiningDate,
                :target_date AS TargetDate,
                DATEDIFF(DAY, e.JoiningDate, :target_date) AS ElapsedServiceDays,
                CAST(COALESCE((SELECT MonthlyRate FROM RateSetting), 150.000) AS DECIMAL(18,3)) AS MonthlyRate,
                CAST(ROUND((CASE WHEN DATEDIFF(DAY, e.JoiningDate, :target_date) < 0 THEN 0 ELSE DATEDIFF(DAY, e.JoiningDate, :target_date) END) * COALESCE((SELECT MonthlyRate FROM RateSetting), 150.000) / 30.4375, 3) AS DECIMAL(18,3)) AS AccruedBalance,
                CAST(COALESCE(s.SeedBalance, 0) AS DECIMAL(18,3)) AS SeedBalance,
                CAST(COALESCE(c.ClaimedBalance, 0) AS DECIMAL(18,3)) AS ClaimedBalance,
                CAST(ROUND(((CASE WHEN DATEDIFF(DAY, e.JoiningDate, :target_date) < 0 THEN 0 ELSE DATEDIFF(DAY, e.JoiningDate, :target_date) END) * COALESCE((SELECT MonthlyRate FROM RateSetting), 150.000) / 30.4375) + COALESCE(s.SeedBalance, 0) - COALESCE(c.ClaimedBalance, 0), 3) AS DECIMAL(18,3)) AS AvailableBalance
            FROM core.Employees e
            LEFT JOIN ApprovedClaims c ON c.EmployeeID = e.EmployeeID
            LEFT JOIN ApprovedSeeds s ON s.EmployeeID = e.EmployeeID
            WHERE e.EmployeeID = :emp_id AND e.IsDeleted = 0
            """
        ),
        {"emp_id": emp_id, "target_date": target_date},
    ).mappings().first()
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Employee not found.")
    return AirfareEntitlementResponse(
        employeeId=int(row["EmployeeID"]),
        employeeCode=row["EmployeeCode"],
        fullName=row["FullName"],
        joiningDate=row["JoiningDate"],
        targetDate=row["TargetDate"],
        elapsedServiceDays=max(0, int(row["ElapsedServiceDays"])),
        monthlyRate=row["MonthlyRate"],
        accruedBalance=row["AccruedBalance"],
        seedBalance=row["SeedBalance"],
        claimedBalance=row["ClaimedBalance"],
        availableBalance=row["AvailableBalance"],
        model="continuous-accrual-elapsed-service-days",
    )


@router.post("/airfare/claims")
def create_airfare_claim(payload: AirfareClaimCreate, session: SessionDep) -> dict[str, Any]:
    result = session.execute(
        text(
            """
            INSERT INTO core.AirfareClaims (EmployeeID, ClaimDate, SectorCode, ClaimType, ClaimAmount, AccruedBalance, ApprovalStatus, Notes)
            OUTPUT INSERTED.ClaimID
            VALUES (:employee_id, :claim_date, :sector_code, :claim_type, :claim_amount, :accrued_balance, N'Submitted', :notes)
            """
        ),
        payload.model_dump(),
    ).scalar_one()
    session.commit()
    return {"claimId": int(result), "status": "Submitted"}


@router.post("/airfare/claims/{claim_id}/approve")
def approve_airfare_claim(claim_id: int, session: SessionDep) -> dict[str, Any]:
    affected = session.execute(
        text("UPDATE core.AirfareClaims SET ApprovalStatus = N'Approved', ApprovedAtUtc = SYSUTCDATETIME() WHERE ClaimID = :claim_id AND ApprovalStatus = N'Submitted'"),
        {"claim_id": claim_id},
    ).rowcount
    session.commit()
    if affected == 0:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Claim not found or not approvable.")
    return {"claimId": claim_id, "status": "Approved"}


@router.post("/loans/amortization")
def generate_amortization(payload: LoanRequestCreate) -> dict[str, Any]:
    monthly_rate = Decimal(payload.annual_interest_rate) / Decimal("1200")
    if monthly_rate == 0:
        emi = payload.principal_amount / Decimal(payload.terms_months)
    else:
        factor = (Decimal("1") + monthly_rate) ** payload.terms_months
        emi = payload.principal_amount * monthly_rate * factor / (factor - Decimal("1"))
    schedule = []
    balance = payload.principal_amount
    for month in range(1, payload.terms_months + 1):
        interest = balance * monthly_rate
        principal = min(balance, emi - interest)
        balance = max(Decimal("0"), balance - principal)
        schedule.append({"installmentNo": month, "principal": round(principal, 3), "interest": round(interest, 3), "emi": round(emi, 3), "balance": round(balance, 3)})
    return {"termsMonths": payload.terms_months, "monthlyInstallment": round(emi, 3), "schedule": schedule}


@router.post("/loans")
def create_loan(payload: LoanRequestCreate, session: SessionDep) -> dict[str, Any]:
    emi = generate_amortization(payload)["monthlyInstallment"]
    loan_id = session.execute(
        text(
            """
            INSERT INTO core.Loans (EmployeeID, PrincipalAmount, TermsMonths, AnnualInterestRate, MonthlyInstallment, OutstandingAmount, DisbursementDate, LoanStatus, Notes)
            OUTPUT INSERTED.LoanID
            VALUES (:employee_id, :principal_amount, :terms_months, :annual_interest_rate, :emi, :principal_amount, :disbursement_date, N'Active', :notes)
            """
        ),
        {**payload.model_dump(), "emi": emi},
    ).scalar_one()
    session.commit()
    return {"loanId": int(loan_id), "monthlyInstallment": emi, "status": "Active"}


@router.post("/seed-evidence")
def create_seed_evidence(payload: SeedEvidenceCreate, session: SessionDep) -> dict[str, Any]:
    seed_id = session.execute(
        text(
            """
            INSERT INTO core.SeedEvidence (EmployeeID, SeedType, EffectiveDate, Amount, SourceReference, VerificationChecksum, ApprovalStatus, ApprovedBy)
            OUTPUT INSERTED.SeedID
            VALUES (:employee_id, :seed_type, :effective_date, :amount, :source_reference, CONVERT(NVARCHAR(64), HASHBYTES('SHA2_256', CONCAT(:employee_id, ':', :seed_type, ':', :effective_date, ':', :amount, ':', :source_reference)), 2), N'Approved', :approved_by)
            """
        ),
        payload.model_dump(),
    ).scalar_one()
    session.commit()
    return {"seedId": int(seed_id), "status": "Approved"}


@router.post("/reports/run")
def run_report(payload: ReportRequest, session: SessionDep) -> dict[str, Any]:
    params = payload.model_dump()
    report_id = session.execute(
        text(
            """
            INSERT INTO core.ReportRuns (ReportCode, DateFrom, DateTo, DepartmentID, EmployeeID, ExportFormat, RunStatus)
            OUTPUT INSERTED.ReportRunID
            VALUES (:report_code, :date_from, :date_to, :department_id, :employee_id, :export_format, N'Completed')
            """
        ),
        params,
    ).scalar_one()
    session.commit()
    return {"reportRunId": int(report_id), "status": "Completed", "parameters": params}


@router.put("/admin/settings/{setting_key}")
def upsert_setting(setting_key: str, payload: SystemSettingUpsert, session: SessionDep) -> dict[str, Any]:
    if setting_key != payload.setting_key:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="setting_key path/body mismatch.")
    session.execute(
        text(
            """
            MERGE core.SystemSettings AS target
            USING (SELECT :setting_key AS SettingKey) AS source
            ON target.SettingKey = source.SettingKey
            WHEN MATCHED THEN UPDATE SET SettingValue = :setting_value, ValueType = :value_type, IsSecret = :is_secret, UpdatedAtUtc = SYSUTCDATETIME()
            WHEN NOT MATCHED THEN INSERT (SettingKey, SettingValue, ValueType, IsSecret) VALUES (:setting_key, :setting_value, :value_type, :is_secret);
            """
        ),
        payload.model_dump(),
    )
    session.commit()
    return {"settingKey": setting_key, "status": "Saved"}


@router.post("/me/requests")
def create_self_service_request(payload: SelfServiceRequestCreate, session: SessionDep) -> dict[str, Any]:
    request_id = session.execute(
        text(
            """
            INSERT INTO core.SelfServiceRequestsV2 (EmployeeID, RequestType, RequestedAtUtc, RequestPayloadJSON, RequestStatus)
            OUTPUT INSERTED.RequestID
            VALUES (:employee_id, :request_type, :requested_at, :payload_json, N'Submitted')
            """
        ),
        {**payload.model_dump(exclude={"payload"}), "payload_json": json_dumps(payload.payload)},
    ).scalar_one()
    session.commit()
    return {"requestId": int(request_id), "status": "Submitted"}


@router.get("/airports", response_model=list[AirportResponse])
def search_airports(session: SessionDep, q: Annotated[str, Query(min_length=1, max_length=80)]) -> list[AirportResponse]:
    like = f"{q.strip()}%"
    contains = f"%{q.strip()}%"
    rows = session.execute(
        text(
            """
            SELECT TOP 20 IATACode, AirportName, CityName, CountryCode, SectorCode,
                   CASE WHEN IATACode = :exact THEN 100 WHEN IATACode LIKE :like THEN 90 WHEN CityName LIKE :like THEN 80 ELSE 50 END AS Score
            FROM core.Airports
            WHERE IsActive = 1 AND (IATACode LIKE :like OR CityName LIKE :contains OR AirportName LIKE :contains OR CountryCode LIKE :like)
            ORDER BY Score DESC, IATACode ASC
            """
        ),
        {"exact": q.strip().upper(), "like": like.upper(), "contains": contains},
    ).all()
    return [AirportResponse(iata_code=r[0], airport_name=r[1], city_name=r[2], country_code=r[3], sector_code=r[4], score=int(r[5])) for r in rows]


@router.post("/attachments/metadata")
def create_document_metadata(file_name: str, employee_id: int, document_type: str, session: SessionDep) -> dict[str, Any]:
    doc_id = session.execute(
        text(
            """
            INSERT INTO core.DocumentMetadata (EmployeeID, DocumentType, FileName, StorageProvider, StoragePath, MimeType, FileSizeBytes, VerificationStatus)
            OUTPUT INSERTED.DocumentID
            VALUES (:employee_id, :document_type, :file_name, N'local', CONCAT(N'documents/', :file_name), N'application/octet-stream', 0, N'Pending')
            """
        ),
        {"employee_id": employee_id, "document_type": document_type, "file_name": file_name},
    ).scalar_one()
    session.commit()
    return {"documentId": int(doc_id), "status": "Pending"}


@router.post("/admin/backups")
def request_backup(session: SessionDep) -> dict[str, Any]:
    backup_id = session.execute(
        text("INSERT INTO core.BackupJobs (JobType, JobStatus, RequestedAtUtc) OUTPUT INSERTED.BackupJobID VALUES (N'DatabaseBackup', N'Queued', SYSUTCDATETIME())")
    ).scalar_one()
    session.commit()
    return {"backupJobId": int(backup_id), "status": "Queued"}


def parse_csv_bytes(content: bytes) -> list[dict[str, str]]:
    text_content = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text_content))
    return [{str(k): str(v or "").strip() for k, v in row.items()} for row in reader]


def parse_import_file(file_name: str, content: bytes) -> list[dict[str, str]]:
    suffix = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else "csv"
    if suffix == "csv":
        return parse_csv_bytes(content)
    if suffix in {"xlsx", "xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(cell or "").strip() for cell in rows[0]]
        parsed: list[dict[str, str]] = []
        for values in rows[1:]:
            parsed.append({headers[index]: "" if value is None else str(value).strip() for index, value in enumerate(values) if index < len(headers)})
        return parsed
    raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Only CSV, XLSX, and XLSM imports are supported.")


def validate_employee_import_row(row: dict[str, str]) -> list[str]:
    required = ["EmployeeCode", "FullName", "FirstName", "LastName", "Gender", "DateOfBirth", "Nationality", "JoiningDate", "Designation", "EmploymentType", "PaymentMode"]
    errors = [f"{field} is required" for field in required if not row.get(field)]
    if row.get("PaymentMode") in {"Bank", "WPS"} and not row.get("IBANAccountNumber"):
        errors.append("IBANAccountNumber is required for Bank/WPS payment.")
    return errors


def validate_employee_preview_row(row: dict[str, str], row_number: int, existing_codes: set[str], seen_codes: set[str]) -> tuple[list[ImportFieldStatus], list[dict[str, Any]], bool]:
    required = ["EmployeeCode", "FullName", "FirstName", "LastName", "Gender", "DateOfBirth", "Nationality", "JoiningDate", "Designation", "EmploymentType", "PaymentMode"]
    allowed = {
        "Gender": {"Male", "Female", "Other", "Undisclosed"},
        "EmploymentType": {"Permanent", "Contract", "Probation", "Temporary", "Intern"},
        "PaymentMode": {"Bank", "Cash", "WPS"},
        "Status": {"Active", "Inactive", "Resigned", "Terminated", "OnLeave", ""},
    }
    date_fields = ["DateOfBirth", "JoiningDate", "ProbationEndDate", "ConfirmationDate", "PassportExpiry", "CivilIDExpiry", "VisaExpiry", "LabourCardExpiry", "ResignationDate", "LastWorkingDay"]
    money_fields = ["BasicSalary", "HousingAllowance", "TransportAllowance", "OtherFixedAllowances"]
    fields: list[ImportFieldStatus] = []
    errors: list[dict[str, Any]] = []

    def add_error(column: str, reason: str) -> None:
        errors.append({"row": row_number, "column": column, "reason": reason})

    employee_code = str(row.get("EmployeeCode") or "").strip().upper()
    duplicate = False
    if employee_code:
        if employee_code in existing_codes or employee_code in seen_codes:
            duplicate = True
            add_error("EmployeeCode", "Duplicate ID")
        seen_codes.add(employee_code)

    for column in sorted(set(required + date_fields + money_fields + list(row.keys()))):
        value = row.get(column)
        status_value = "VALID"
        reason = None
        if column in required and not str(value or "").strip():
            status_value = "ERROR"
            reason = "Required"
            add_error(column, "Required")
        elif column in date_fields and str(value or "").strip():
            try:
                date.fromisoformat(str(value).strip()[:10])
            except ValueError:
                status_value = "ERROR"
                reason = "Invalid Format"
                add_error(column, "Invalid Format")
        elif column in money_fields and str(value or "").strip():
            try:
                if Decimal(str(value).strip()) < 0:
                    raise ValueError
            except Exception:
                status_value = "ERROR"
                reason = "Invalid non-negative amount"
                add_error(column, "Invalid non-negative amount")
        elif column in allowed and str(value or "").strip() not in allowed[column]:
            status_value = "ERROR"
            reason = "Invalid value"
            add_error(column, "Invalid value")
        elif column == "EmployeeCode" and duplicate:
            status_value = "ERROR"
            reason = "Duplicate ID"
        elif column not in EMPLOYEE_IMPORT_COLUMNS:
            status_value = "WARNING"
            reason = "Ignored column"
        fields.append(ImportFieldStatus(column=column, value=value, status=status_value, reason=reason))
    return fields, errors, duplicate


def normalize_employee_import_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized = {column: row.get(column) or None for column in EMPLOYEE_IMPORT_COLUMNS}
    normalized["Status"] = normalized.get("Status") or "Active"
    normalized["BasicSalary"] = str(normalized.get("BasicSalary") or "0.000")
    normalized["HousingAllowance"] = str(normalized.get("HousingAllowance") or "0.000")
    normalized["TransportAllowance"] = str(normalized.get("TransportAllowance") or "0.000")
    normalized["OtherFixedAllowances"] = str(normalized.get("OtherFixedAllowances") or "0.000")
    normalized["RehireEligible"] = 1 if str(normalized.get("RehireEligible") or "true").lower() in {"1", "true", "yes", "y"} else 0
    return normalized


EMPLOYEE_IMPORT_COLUMNS = [
    "EmployeeCode", "PunchMachineID", "FullName", "FirstName", "MiddleName", "LastName", "PassportName", "Gender", "DateOfBirth", "Nationality", "Religion", "MaritalStatus",
    "JoiningDate", "ProbationEndDate", "ConfirmationDate", "Designation", "GradeLevel", "EmploymentType", "Status",
    "PersonalEmail", "WorkEmail", "MobileNumber", "EmergencyContactName", "EmergencyContactPhone", "EmergencyContactRelationship", "LocalAddress", "HomeCountryAddress",
    "PassportNumber", "PassportExpiry", "CivilID", "CivilIDExpiry", "VisaNumber", "VisaType", "VisaExpiry", "LabourCardNumber", "LabourCardExpiry",
    "BasicSalary", "HousingAllowance", "TransportAllowance", "OtherFixedAllowances", "PaymentMode", "BankName", "IBANAccountNumber", "SwiftCode",
    "ResignationDate", "LastWorkingDay", "ReasonForLeaving", "RehireEligible",
]


EMPLOYEE_IMPORT_INSERT_SQL = text(
    """
    INSERT INTO core.Employees (
        EmployeeCode, PunchMachineID, FullName, FirstName, MiddleName, LastName, PassportName, Gender, DateOfBirth, Nationality, Religion, MaritalStatus,
        JoiningDate, ProbationEndDate, ConfirmationDate, Designation, GradeLevel, EmploymentType, Status,
        PersonalEmail, WorkEmail, MobileNumber, EmergencyContactName, EmergencyContactPhone, EmergencyContactRelationship, LocalAddress, HomeCountryAddress,
        PassportNumber, PassportExpiry, CivilID, CivilIDExpiry, VisaNumber, VisaType, VisaExpiry, LabourCardNumber, LabourCardExpiry,
        BasicSalary, HousingAllowance, TransportAllowance, OtherFixedAllowances, PaymentMode, BankName, IBANAccountNumber, SwiftCode,
        ResignationDate, LastWorkingDay, ReasonForLeaving, RehireEligible
    )
    VALUES (
        :EmployeeCode, :PunchMachineID, :FullName, :FirstName, :MiddleName, :LastName, :PassportName, :Gender, :DateOfBirth, :Nationality, :Religion, :MaritalStatus,
        :JoiningDate, :ProbationEndDate, :ConfirmationDate, :Designation, :GradeLevel, :EmploymentType, :Status,
        :PersonalEmail, :WorkEmail, :MobileNumber, :EmergencyContactName, :EmergencyContactPhone, :EmergencyContactRelationship, :LocalAddress, :HomeCountryAddress,
        :PassportNumber, :PassportExpiry, :CivilID, :CivilIDExpiry, :VisaNumber, :VisaType, :VisaExpiry, :LabourCardNumber, :LabourCardExpiry,
        :BasicSalary, :HousingAllowance, :TransportAllowance, :OtherFixedAllowances, :PaymentMode, :BankName, :IBANAccountNumber, :SwiftCode,
        :ResignationDate, :LastWorkingDay, :ReasonForLeaving, :RehireEligible
    )
    """
)


def json_dumps(value: dict[str, Any]) -> str:
    return json.dumps(value, separators=(",", ":"), default=str)

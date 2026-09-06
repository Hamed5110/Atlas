"""Report export integration tests."""

from __future__ import annotations

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.helpers import create_employee, create_ticket

pytestmark = pytest.mark.integration

REPORT_NAMES = (
    "employee-master",
    "opening-balances",
    "entitlements",
    "ticket-register",
    "loan-outstanding",
    "loan-statement",
    "excess-recovery",
    "liability-projections",
    "airfare-payable",
    "airfare-payable-summary",
    "airfare-payable-exceptions",
)


@pytest.mark.parametrize("report_name", REPORT_NAMES)
def test_report_detail_structure(
    client: TestClient, admin_headers: dict[str, str], report_name: str
) -> None:
    response = client.get(f"/v1/reports/detail/{report_name}", headers=admin_headers)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["report"] == report_name
    assert isinstance(body["columns"], list)
    assert isinstance(body["rows"], list)
    assert body["count"] == len(body["rows"])


@pytest.mark.parametrize("report_name", REPORT_NAMES)
def test_report_data_tile(client: TestClient, admin_headers: dict[str, str], report_name: str) -> None:
    response = client.get(f"/v1/reports/data/{report_name}", headers=admin_headers)
    assert response.status_code == 200
    body = response.json()
    assert body["report"] == report_name
    assert "value" in body


def test_excel_export_sheet_and_headers(client: TestClient, admin_headers: dict[str, str]) -> None:
    employee = create_employee(client, admin_headers, full_name="=HYPERLINK(\"evil\")")
    create_ticket(client, admin_headers, str(employee["id"]))
    response = client.get(
        "/v1/reports/export/ticket-register.xlsx",
        headers=admin_headers,
    )
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    assert sheet.max_row >= 1
    headers = [cell.value for cell in sheet[1]]
    assert "Travel date" in headers


def test_pdf_export_is_valid_pdf(client: TestClient, admin_headers: dict[str, str]) -> None:
    response = client.get("/v1/reports/export/employee-master.pdf", headers=admin_headers)
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert b"Employee" in response.content or b"employee" in response.content.lower()


def test_formula_injection_sanitized_in_report_excel(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    create_employee(client, admin_headers, full_name="=1+1", code="FXRPT1")
    response = client.get("/v1/reports/export/employee-master.xlsx", headers=admin_headers)
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None
    values = [str(cell.value or "") for row in sheet.iter_rows(min_row=2) for cell in row]
    assert any(value.startswith("'=") for value in values)


def test_excess_pdf_alias_endpoint(client: TestClient, admin_headers: dict[str, str]) -> None:
    response = client.get("/v1/reports/excess.pdf", headers=admin_headers)
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")

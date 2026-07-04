from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ATLAS_Airfare_HCM_Support_Guide_2026-06-28.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
BORDER = "B7C9D9"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, name="Calibri", size=11, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("ATLAS Airfare HCM Support Guide - Verified 2026-06-28")
    set_font(run, size=9, color=MUTED)


def add_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("ATLAS Airfare HCM Support Guide")
    set_font(run, size=22, color=BLUE, bold=True)

    p = doc.add_paragraph()
    run = p.add_run("Operator help, testing status, and issue-response reference")
    set_font(run, size=12, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    set_table_width(table, [2400, 6960])
    set_table_borders(table)
    rows = [
        ("Verified date", "2026-06-28"),
        ("Primary URL", "http://<server-name>/"),
        ("Fallback URL", "http://192.168.15.208/"),
        ("Testing result", "Passed acceptance, full system, frontend, SQL, and live UI checks"),
    ]
    for idx, (label, value) in enumerate(rows):
        table.cell(idx, 0).text = label
        table.cell(idx, 1).text = value
        set_cell_fill(table.cell(idx, 0), HEADER_FILL)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.add_run(item)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.add_run(item)


def add_section(doc, title, items):
    doc.add_heading(title, level=1)
    add_bullets(doc, items)


def add_status_table(doc):
    doc.add_heading("Verified Test Coverage", level=1)
    table = doc.add_table(rows=1, cols=3)
    set_table_width(table, [2100, 2100, 5160])
    set_table_borders(table)
    headers = ["Area", "Result", "Evidence"]
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        set_cell_fill(cell, HEADER_FILL)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_font(run, bold=True, color=DARK_BLUE)

    rows = [
        ("Year End", "Passed", "Preview, dry-run close, and temporary future-year close carry-forward were verified."),
        ("Loans", "Passed", "Loan SQL, register, EMI preview/run, restructure, deferment, settlement, and history were verified."),
        ("Reports", "Passed", "Employee Master, Airfare Payable, Year Summary, dashboard, export/source checks, and report viewer were verified."),
        ("Companies", "Passed", "Company list, backup list, temporary company create/update/backup/restore were verified."),
        ("Frontend", "Passed", "Production build, live login, main screens, and console error check passed."),
    ]
    for area, result, evidence in rows:
        cells = table.add_row().cells
        cells[0].text = area
        cells[1].text = result
        cells[2].text = evidence
        set_cell_fill(cells[1], LIGHT_FILL)


def add_troubleshooting(doc):
    doc.add_heading("Quick Troubleshooting", level=1)
    table = doc.add_table(rows=1, cols=3)
    set_table_width(table, [2400, 2760, 4200])
    set_table_borders(table)
    headers = ["Symptom", "First Check", "Action"]
    for i, header in enumerate(headers):
        table.cell(0, i).text = header
        set_cell_fill(table.cell(0, i), HEADER_FILL)

    rows = [
        ("Other PC cannot open name", "Try http://192.168.15.208/", "Run Open-ATLAS-From-Other-PC.bat on the other PC."),
        ("Login opens but data fails", "Open /api/health", "Refresh once, then restart with Start-ATLAS-LAN.ps1 if health fails."),
        ("Year-end warning", "Review pending loans", "Preview again after loan review; close only after balances are checked."),
        ("Report value question", "Compare same employee across screens", "Check Employee Master, Airfare Allocation, and Airfare Payable before editing data."),
    ]
    for symptom, check, action in rows:
        cells = table.add_row().cells
        cells[0].text = symptom
        cells[1].text = check
        cells[2].text = action


def build():
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("Daily Operating Sequence", level=1)
    add_steps(doc, [
        "Open ATLAS and sign in.",
        "Confirm the selected company.",
        "Review Overview totals for airfare payable, opening balance, and loan exposure.",
        "Maintain Employee Master before allocations.",
        "Enter or import Opening Balance before ticket processing.",
        "Process Airfare Allocation with ticket amount, payment option, and approval where required.",
        "Review Loans for EMI, deferment, restructure, or settlement.",
        "Use Reports for review, export, and print.",
        "Run Year End preview before any close.",
    ])

    add_section(doc, "Company Administration", [
        "Use Companies to create and maintain company records.",
        "Use Create Backup before major company or database work.",
        "Use Restore only with a selected backup file and administrator approval.",
        "Do not delete the main ATLAS company.",
    ])
    add_section(doc, "Employee Master and Opening Balance", [
        "Keep employee status accurate; only active employees are eligible.",
        "Maintain opening balances in the Opening Balance screen.",
        "Review Excel import preview rows before confirming.",
        "Opening amount follows the active maximum payout policy.",
    ])
    add_section(doc, "Airfare Allocation", [
        "Select employee, allocation date, ticket amount, and payment mode.",
        "Manager approval is required for second tickets and all loan tickets.",
        "Attach ticket or invoice documents where required.",
        "Excess can be employee self-pay or loan depending on approval and selected payment option.",
    ])
    add_section(doc, "Loans", [
        "Manual loans can be created from the Loans screen.",
        "Airfare excess loans are created from Airfare Allocation after manager approval.",
        "Use EMI preview before running monthly deductions.",
        "Review loan history after restructure, deferment, EMI run, or settlement.",
    ])
    add_section(doc, "Year End", [
        "Preview year end before closing.",
        "Review pending loans and closing balance.",
        "Close year only after checking employee balances and reports.",
        "Closing balance is carried forward as next year opening balance.",
    ])
    add_section(doc, "Reports", [
        "Use Reports for dashboard, employee master, airfare payable, allocation register, loan register, company register, and year-end preview.",
        "Use date filters before exporting or printing.",
        "Airfare Payable is the control report for payable entitlement and status review.",
    ])
    add_status_table(doc)
    add_troubleshooting(doc)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()

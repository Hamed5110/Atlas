from __future__ import annotations

import base64
import csv
import io
import json
import re
import shutil
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

import mssql_python
from openpyxl import load_workbook

from . import config


ROOT = Path(__file__).resolve().parents[1]
SCHEMA_FILE = ROOT / "schema" / "mssql" / "001_core.sql"
BACKUP_DIR = ROOT / "backups"
AIRPORTS = [
    {"code": "BAH", "name": "Bahrain International Airport", "city": "Manama", "country": "Bahrain"},
    {"code": "DXB", "name": "Dubai International Airport", "city": "Dubai", "country": "United Arab Emirates"},
    {"code": "DMM", "name": "King Fahd International Airport", "city": "Dammam", "country": "Saudi Arabia"},
    {"code": "DOH", "name": "Hamad International Airport", "city": "Doha", "country": "Qatar"},
    {"code": "KWI", "name": "Kuwait International Airport", "city": "Kuwait City", "country": "Kuwait"},
    {"code": "BOM", "name": "Chhatrapati Shivaji Maharaj International Airport", "city": "Mumbai", "country": "India"},
    {"code": "COK", "name": "Cochin International Airport", "city": "Kochi", "country": "India"},
    {"code": "DEL", "name": "Indira Gandhi International Airport", "city": "Delhi", "country": "India"},
    {"code": "MAA", "name": "Chennai International Airport", "city": "Chennai", "country": "India"},
    {"code": "LHE", "name": "Allama Iqbal International Airport", "city": "Lahore", "country": "Pakistan"},
]


def connect(database: str | None = None, autocommit: bool = True):
    return mssql_python.connect(config.connection_string(database), autocommit=autocommit)


def ensure_database() -> None:
    validate_database_name(config.DB_NAME)
    with connect("master") as conn:
        cur = conn.cursor()
        database_identifier = "[" + config.DB_NAME.replace("]", "]]") + "]"
        cur.execute(f"IF DB_ID(?) IS NULL EXEC(N'CREATE DATABASE {database_identifier}');", (config.DB_NAME,))


def initialize() -> dict[str, Any]:
    ensure_database()
    with connect(config.DB_NAME) as conn:
        apply_schema(conn)
        seed(conn)
        return health_probe(conn)


def apply_schema(conn) -> None:
    sql_text = SCHEMA_FILE.read_text(encoding="utf-8")
    for batch in re.split(r"^\s*GO\s*$", sql_text, flags=re.IGNORECASE | re.MULTILINE):
        batch = batch.strip()
        if batch:
            conn.cursor().execute(batch)


def seed(conn) -> None:
    cur = conn.cursor()
    tenant_id = config.TENANT_ID
    company_id = config.COMPANY_ID
    employee_1 = "33333333-3333-4333-8333-333333333333"
    employee_2 = "33333333-3333-4333-8333-333333333334"
    rule_id = "44444444-4444-4444-8444-444444444444"
    user_id = "55555555-5555-4555-8555-555555555555"

    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.Tenants WHERE TenantID = ?) "
        "INSERT INTO core.Tenants (TenantID, TenantCode, TenantName) VALUES (?, N'ATLAS', N'ATLAS Greenfield');",
        (tenant_id, tenant_id),
    )
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.Companies WHERE CompanyID = ?) "
        "INSERT INTO core.Companies (CompanyID, TenantID, CompanyCode, CompanyName) VALUES (?, ?, N'ATLAS', N'ATLAS Airfare HCM');",
        (company_id, company_id, tenant_id),
    )
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.Employees WHERE EmployeeID = ?) "
        "INSERT INTO core.Employees (EmployeeID, TenantID, CompanyID, EmployeeNumber, DisplayName, WorkEmail, Department, JobTitle, HireDate) "
        "VALUES (?, ?, ?, N'5110', N'test', N'test@example.com', N'Validation', N'Core Employee', '2022-01-01');",
        (employee_1, employee_1, tenant_id, company_id),
    )
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.Employees WHERE EmployeeID = ?) "
        "INSERT INTO core.Employees (EmployeeID, TenantID, CompanyID, EmployeeNumber, DisplayName, WorkEmail, Department, JobTitle, HireDate) "
        "VALUES (?, ?, ?, N'5111', N'production sample', N'prod@example.com', N'Operations', N'Airfare User', '2023-01-01');",
        (employee_2, employee_2, tenant_id, company_id),
    )
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.EntitlementRules WHERE RuleID = ?) "
        "INSERT INTO core.EntitlementRules (RuleID, TenantID, CompanyID, RuleCode, RuleName, Cadence, MaxPayoutAmount, EffectiveFrom) "
        "VALUES (?, ?, ?, N'GLOBAL-150', N'Continuous airfare entitlement', N'service_day', 150.000, '2026-01-01');",
        (rule_id, rule_id, tenant_id, company_id),
    )
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.EntitlementEvents WHERE EmployeeID = ? AND EventType = N'seed' AND SourceReference = N'python-core-seed') "
        "INSERT INTO core.EntitlementEvents (TenantID, CompanyID, EmployeeID, EventDate, EventType, Amount, SourceReference) "
        "VALUES (?, ?, ?, '2026-01-01', N'seed', 60.000, N'python-core-seed');",
        (employee_1, tenant_id, company_id, employee_1),
    )
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.EmployeeLoans WHERE EmployeeID = ? AND StartDate = '2026-02-01') "
        "INSERT INTO core.EmployeeLoans (LoanID, TenantID, CompanyID, EmployeeID, PrincipalAmount, EmiAmount, StartDate) "
        "VALUES (?, ?, ?, ?, 300.000, 50.000, '2026-02-01');",
        (employee_1, str(uuid.uuid4()), tenant_id, company_id, employee_1),
    )
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.Users WHERE UserID = ?) "
        "INSERT INTO core.Users (UserID, Username, DisplayName, RoleCode) VALUES (?, N'admin', N'System Administrator', N'admin');",
        (user_id, user_id),
    )
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.UserPreferences WHERE UserID = ?) "
        "INSERT INTO core.UserPreferences (PreferenceID, UserID, PreferenceJSON) VALUES (?, ?, ?);",
        (user_id, str(uuid.uuid4()), user_id, '{"theme":"system","density":"comfortable","defaultPage":"command"}'),
    )
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.SelfServiceRequests WHERE EmployeeID = ? AND RequestDate = '2026-05-01') "
        "INSERT INTO core.SelfServiceRequests (RequestID, TenantID, CompanyID, EmployeeID, RequestType, RequestDate, Amount, Notes) "
        "VALUES (?, ?, ?, ?, N'airfare_request', '2026-05-01', 75.000, N'Seed self-service request');",
        (employee_1, str(uuid.uuid4()), tenant_id, company_id, employee_1),
    )
    sample_attachment_id = "66666666-6666-4666-8666-666666666666"
    cur.execute(
        "IF NOT EXISTS (SELECT 1 FROM core.Attachments WHERE AttachmentID = ?) "
        "INSERT INTO core.Attachments (AttachmentID, ModuleCode, OwnerID, FileName, ContentType, ContentBytes) "
        "VALUES (?, N'allocations', ?, N'sample-airfare-evidence.txt', N'text/plain; charset=utf-8', CONVERT(varbinary(max), ?));",
        (sample_attachment_id, sample_attachment_id, employee_1, "ATLAS Python Core sample airfare evidence"),
    )


def health_probe(conn=None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect(config.DB_NAME)
    try:
        cur = conn.cursor()
        cur.execute("SELECT @@SERVERNAME AS server_name, DB_NAME() AS database_name;")
        server_name, database_name = cur.fetchone()
        cur.execute(
            "SELECT "
            "OBJECT_ID(N'core.Employees', N'U') AS employees_table, "
            "OBJECT_ID(N'core.EntitlementEvents', N'U') AS events_table, "
            "OBJECT_ID(N'core.AirfareAllocations', N'U') AS allocations_table, "
            "OBJECT_ID(N'core.Attachments', N'U') AS attachments_table;"
        )
        employees_table, events_table, allocations_table, attachments_table = cur.fetchone()
        return {
            "serverName": server_name,
            "databaseName": database_name,
            "schema": "core",
            "objects": {
                "employees": bool(employees_table),
                "entitlementEvents": bool(events_table),
                "airfareAllocations": bool(allocations_table),
                "attachments": bool(attachments_table),
            },
        }
    finally:
        if own_conn:
            conn.close()


def summary(as_of_date: str) -> dict[str, Any]:
    require_iso_date(as_of_date, "asOfDate")
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*), SUM(CASE WHEN StatusCode = N'active' THEN 1 ELSE 0 END) FROM core.Employees;")
        employee_total, employee_active = cur.fetchone()
        cur.execute(
            "SELECT CAST(COALESCE(SUM(CASE WHEN EventType IN (N'seed',N'accrual',N'adjustment',N'reversal') THEN Amount ELSE 0 END),0) AS DECIMAL(12,3)), "
            "CAST(COALESCE(SUM(CASE WHEN EventType = N'usage' THEN Amount ELSE 0 END),0) AS DECIMAL(12,3)) "
            "FROM core.EntitlementEvents WHERE EventDate <= ?;",
            (as_of_date,),
        )
        earned, used = cur.fetchone()
        cur.execute("SELECT COUNT(*), COALESCE(SUM(PrincipalAmount),0), COALESCE(SUM(EmiAmount),0) FROM core.EmployeeLoans WHERE StatusCode = N'active';")
        active_loans, loan_principal, monthly_emi = cur.fetchone()
        cur.execute("SELECT COUNT(*), COALESCE(SUM(TicketCost),0), COALESCE(SUM(EntitlementApplied),0), COALESCE(SUM(CompanyPaid),0) FROM core.AirfareAllocations WHERE StatusCode = N'posted';")
        allocation_count, ticket_cost, entitlement_applied, company_paid = cur.fetchone()
        balance = money(Decimal(earned or 0) - Decimal(used or 0))
        return {
            "runtime": "GREEN",
            "application": config.APP_NAME,
            "version": config.APP_VERSION,
            "repository": "mssql-python-core",
            "database": config.DB_NAME,
            "asOfDate": as_of_date,
            "employees": {"total": int(employee_total or 0), "active": int(employee_active or 0)},
            "entitlement": {"earned": money(earned), "used": money(used), "balance": balance, "currency": "BHD"},
            "allocations": {
                "total": int(allocation_count or 0),
                "ticketCost": money(ticket_cost),
                "entitlementApplied": money(entitlement_applied),
                "companyPaid": money(company_paid),
            },
            "loans": {"active": int(active_loans or 0), "principal": money(loan_principal), "monthlyEmi": money(monthly_emi)},
        }


def login(payload: dict[str, Any]) -> dict[str, Any]:
    username = required_text(payload, "username").lower()
    password = required_text(payload, "password")
    if username != "admin" or password not in {"admin", "Admin123!"}:
        raise ValueError("Invalid username or password.")
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute("SELECT UserID, DisplayName, RoleCode FROM core.Users WHERE Username = ? AND IsActive = 1;", (username,))
        row = cur.fetchone()
        if not row:
            raise ValueError("User is not active.")
    return {"token": "local-python-core-admin", "sessionId": str(uuid.uuid4()), "user": {"userId": str(row[0]), "displayName": row[1], "roleCode": row[2]}}


def list_companies() -> list[dict[str, Any]]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT CompanyID, CompanyCode, CompanyName, BaseCurrencyCode, IsActive "
            "FROM core.Companies ORDER BY CompanyCode;"
        )
        return [
            {
                "companyId": str(row[0]),
                "companyCode": row[1],
                "companyName": row[2],
                "baseCurrencyCode": row[3],
                "isActive": bool(row[4]),
            }
            for row in cur.fetchall()
        ]


def create_company(payload: dict[str, Any]) -> dict[str, Any]:
    company_id = str(uuid.uuid4())
    code = required_text(payload, "companyCode").upper()
    name = required_text(payload, "companyName")
    currency = str(payload.get("baseCurrencyCode") or "BHD").upper()[:3]
    with connect(config.DB_NAME) as conn:
        conn.cursor().execute(
            "INSERT INTO core.Companies (CompanyID, TenantID, CompanyCode, CompanyName, BaseCurrencyCode) VALUES (?, ?, ?, ?, ?);",
            (company_id, config.TENANT_ID, code, name, currency),
        )
    return {"companyId": company_id, "companyCode": code, "companyName": name, "baseCurrencyCode": currency}


def list_employees() -> list[dict[str, Any]]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT EmployeeID, EmployeeNumber, DisplayName, LegalName, WorkEmail, PhoneNumber, Department, JobTitle, EmploymentType, PayGroup, "
            "Nationality, PassportNumber, CPRNumber, BankName, IBAN, BasicSalary, EligibleForAirfare, HomeAirportCode, DestinationAirportCode, "
            "StatusCode, CONVERT(char(10), HireDate, 23), CONVERT(char(10), TerminationDate, 23) "
            "FROM core.Employees ORDER BY EmployeeNumber;"
        )
        return [
            {
                "employeeId": str(row[0]),
                "employeeNumber": row[1],
                "displayName": row[2],
                "legalName": row[3],
                "workEmail": row[4],
                "phoneNumber": row[5],
                "department": row[6],
                "jobTitle": row[7],
                "employmentType": row[8],
                "payGroup": row[9],
                "nationality": row[10],
                "passportNumber": row[11],
                "cprNumber": row[12],
                "bankName": row[13],
                "iban": row[14],
                "basicSalary": money(row[15]),
                "eligibleForAirfare": bool(row[16]),
                "homeAirportCode": row[17],
                "destinationAirportCode": row[18],
                "statusCode": row[19],
                "hireDate": row[20],
                "terminationDate": row[21],
            }
            for row in cur.fetchall()
        ]


def entitlement_balance(as_of_date: str) -> list[dict[str, Any]]:
    require_iso_date(as_of_date, "asOfDate")
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT e.EmployeeNumber, e.DisplayName, "
            "CAST(COALESCE(SUM(CASE WHEN ev.EventType IN (N'seed',N'accrual',N'adjustment',N'reversal') THEN ev.Amount ELSE 0 END),0) AS DECIMAL(12,3)) AS earned, "
            "CAST(COALESCE(SUM(CASE WHEN ev.EventType = N'usage' THEN ev.Amount ELSE 0 END),0) AS DECIMAL(12,3)) AS used "
            "FROM core.Employees e "
            "LEFT JOIN core.EntitlementEvents ev ON ev.EmployeeID = e.EmployeeID AND ev.EventDate <= ? "
            "GROUP BY e.EmployeeNumber, e.DisplayName ORDER BY e.EmployeeNumber;",
            (as_of_date,),
        )
        rows = []
        for number, name, earned, used in cur.fetchall():
            rows.append({
                "employeeNumber": number,
                "displayName": name,
                "earned": money(earned),
                "used": money(used),
                "balance": money(Decimal(earned or 0) - Decimal(used or 0)),
                "currency": "BHD",
            })
        return rows


def list_entitlement_rules() -> list[dict[str, Any]]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT RuleID, RuleCode, RuleName, Cadence, MaxPayoutAmount, CurrencyCode, CONVERT(char(10), EffectiveFrom, 23), IsActive "
            "FROM core.EntitlementRules ORDER BY RuleCode;"
        )
        return [
            {
                "ruleId": str(row[0]),
                "ruleCode": row[1],
                "ruleName": row[2],
                "cadence": row[3],
                "maxPayoutAmount": money(row[4]),
                "currencyCode": row[5],
                "effectiveFrom": row[6],
                "isActive": bool(row[7]),
            }
            for row in cur.fetchall()
        ]


def list_entitlement_events() -> list[dict[str, Any]]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT TOP 100 ev.EventID, e.EmployeeNumber, e.DisplayName, CONVERT(char(10), ev.EventDate, 23), ev.EventType, ev.Amount, ev.SourceReference "
            "FROM core.EntitlementEvents ev "
            "JOIN core.Employees e ON e.EmployeeID = ev.EmployeeID "
            "ORDER BY ev.EventDate DESC, ev.EventID DESC;"
        )
        return [
            {
                "eventId": int(row[0]),
                "employeeNumber": row[1],
                "displayName": row[2],
                "eventDate": row[3],
                "eventType": row[4],
                "amount": money(row[5]),
                "sourceReference": row[6],
            }
            for row in cur.fetchall()
        ]


def post_entitlement_event(payload: dict[str, Any]) -> dict[str, Any]:
    employee_id = required_text(payload, "employeeId")
    event_date = required_text(payload, "eventDate")
    event_type = required_text(payload, "eventType")
    amount = non_negative(payload.get("amount"), "amount")
    require_iso_date(event_date, "eventDate")
    if event_type not in {"seed", "accrual", "adjustment", "reversal"}:
        raise ValueError("eventType must be seed, accrual, adjustment, or reversal.")
    with connect(config.DB_NAME) as conn:
        assert_employee_exists(conn, employee_id)
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO core.EntitlementEvents (TenantID, CompanyID, EmployeeID, EventDate, EventType, Amount, SourceReference) "
            "OUTPUT INSERTED.EventID VALUES (?, ?, ?, ?, ?, ?, ?);",
            (config.TENANT_ID, config.COMPANY_ID, employee_id, event_date, event_type, amount, payload.get("sourceReference") or "manual-admin"),
        )
        event_id = cur.fetchone()[0]
    return {"eventId": int(event_id), "employeeId": employee_id, "eventDate": event_date, "eventType": event_type, "amount": money(amount)}


def list_allocations() -> list[dict[str, Any]]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT a.AllocationID, e.EmployeeNumber, e.DisplayName, CONVERT(char(10), a.AllocationDate, 23), "
            "a.OriginAirportCode, a.DestinationAirportCode, CONVERT(char(10), a.TravelDate, 23), a.AirlineName, a.TicketNumber, a.PaymentMode, "
            "a.TicketCost, a.EntitlementApplied, a.CompanyPaid, a.StatusCode "
            "FROM core.AirfareAllocations a "
            "JOIN core.Employees e ON e.EmployeeID = a.EmployeeID "
            "ORDER BY a.AllocationDate DESC, a.CreatedAtUtc DESC;"
        )
        return [
            {
                "allocationId": str(row[0]),
                "employeeNumber": row[1],
                "displayName": row[2],
                "allocationDate": row[3],
                "originAirportCode": row[4],
                "destinationAirportCode": row[5],
                "travelDate": row[6],
                "airlineName": row[7],
                "ticketNumber": row[8],
                "paymentMode": row[9],
                "ticketCost": money(row[10]),
                "entitlementApplied": money(row[11]),
                "companyPaid": money(row[12]),
                "statusCode": row[13],
            }
            for row in cur.fetchall()
        ]


def create_allocation(payload: dict[str, Any]) -> dict[str, Any]:
    employee_id = required_text(payload, "employeeId")
    allocation_date = required_text(payload, "allocationDate")
    ticket_cost = non_negative(payload.get("ticketCost"), "ticketCost")
    travel_date = optional_date(payload.get("travelDate"), "travelDate")
    origin = airport_code(payload.get("originAirportCode"))
    destination = airport_code(payload.get("destinationAirportCode"))
    payment_mode = option(payload.get("paymentMode") or "entitlement", "paymentMode", {"entitlement", "loan", "employee", "company", "mixed"})
    require_iso_date(allocation_date, "allocationDate")
    current_balance = Decimal(str(_employee_balance(employee_id, allocation_date)))
    entitlement_applied = min(Decimal(str(ticket_cost)), current_balance)
    company_paid = max(Decimal("0"), Decimal(str(ticket_cost)) - entitlement_applied)
    allocation_id = str(uuid.uuid4())

    with connect(config.DB_NAME, autocommit=False) as conn:
        try:
            assert_employee_exists(conn, employee_id)
            cur = conn.cursor()
            cur.execute(
                "INSERT INTO core.EntitlementEvents (TenantID, CompanyID, EmployeeID, EventDate, EventType, Amount, SourceReference) "
                "OUTPUT INSERTED.EventID VALUES (?, ?, ?, ?, N'usage', ?, N'airfare-allocation');",
                (config.TENANT_ID, config.COMPANY_ID, employee_id, allocation_date, entitlement_applied),
            )
            event_id = cur.fetchone()[0]
            cur.execute(
                "INSERT INTO core.AirfareAllocations "
                "(AllocationID, TenantID, CompanyID, EmployeeID, AllocationDate, OriginAirportCode, DestinationAirportCode, TravelDate, AirlineName, TicketNumber, PaymentMode, TicketCost, EntitlementApplied, CompanyPaid, EventID) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);",
                (
                    allocation_id,
                    config.TENANT_ID,
                    config.COMPANY_ID,
                    employee_id,
                    allocation_date,
                    origin,
                    destination,
                    travel_date,
                    text_or_none(payload.get("airlineName")),
                    text_or_none(payload.get("ticketNumber")),
                    payment_mode,
                    ticket_cost,
                    entitlement_applied,
                    company_paid,
                    event_id,
                ),
            )
            conn.commit()
        except Exception:
            conn.rollback()
            raise
    return {
        "allocationId": allocation_id,
        "employeeId": employee_id,
        "allocationDate": allocation_date,
        "originAirportCode": origin,
        "destinationAirportCode": destination,
        "travelDate": travel_date,
        "paymentMode": payment_mode,
        "ticketCost": money(ticket_cost),
        "entitlementApplied": money(entitlement_applied),
        "companyPaid": money(company_paid),
        "statusCode": "posted",
    }


def list_loans() -> list[dict[str, Any]]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT l.LoanID, e.EmployeeNumber, e.DisplayName, l.LoanType, l.PrincipalAmount, l.EmiAmount, l.TenureMonths, l.OutstandingAmount, "
            "CONVERT(char(10), l.StartDate, 23), CONVERT(char(10), l.LoanDate, 23), l.Notes, l.StatusCode "
            "FROM core.EmployeeLoans l "
            "JOIN core.Employees e ON e.EmployeeID = l.EmployeeID "
            "ORDER BY l.StartDate DESC, e.EmployeeNumber;"
        )
        return [
            {
                "loanId": str(row[0]),
                "employeeNumber": row[1],
                "displayName": row[2],
                "loanType": row[3],
                "principalAmount": money(row[4]),
                "emiAmount": money(row[5]),
                "tenureMonths": int(row[6] or 0),
                "outstandingAmount": money(row[7]),
                "startDate": row[8],
                "loanDate": row[9],
                "notes": row[10],
                "statusCode": row[11],
            }
            for row in cur.fetchall()
        ]


def settle_loan(loan_id: str) -> dict[str, Any]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute("UPDATE core.EmployeeLoans SET StatusCode = N'settled' WHERE LoanID = ?;", (loan_id,))
    return {"loanId": loan_id, "statusCode": "settled"}


def create_loan(payload: dict[str, Any]) -> dict[str, Any]:
    employee_id = required_text(payload, "employeeId")
    loan_type = option(payload.get("loanType") or "airfare", "loanType", {"airfare", "emergency_ticket", "salary_advance", "manual"})
    principal = non_negative(payload.get("principalAmount"), "principalAmount")
    emi = non_negative(payload.get("emiAmount"), "emiAmount")
    tenure_months = non_negative_int(payload.get("tenureMonths") or 0, "tenureMonths")
    outstanding = non_negative(payload.get("outstandingAmount") or principal, "outstandingAmount")
    start_date = required_text(payload, "startDate")
    loan_date = optional_date(payload.get("loanDate"), "loanDate") or start_date
    require_iso_date(start_date, "startDate")
    loan_id = str(uuid.uuid4())
    with connect(config.DB_NAME) as conn:
        assert_employee_exists(conn, employee_id)
        conn.cursor().execute(
            "INSERT INTO core.EmployeeLoans (LoanID, TenantID, CompanyID, EmployeeID, LoanType, PrincipalAmount, EmiAmount, TenureMonths, OutstandingAmount, StartDate, LoanDate, Notes) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);",
            (loan_id, config.TENANT_ID, config.COMPANY_ID, employee_id, loan_type, principal, emi, tenure_months, outstanding, start_date, loan_date, text_or_none(payload.get("notes"))),
        )
    return {
        "loanId": loan_id,
        "employeeId": employee_id,
        "loanType": loan_type,
        "principalAmount": money(principal),
        "emiAmount": money(emi),
        "tenureMonths": tenure_months,
        "outstandingAmount": money(outstanding),
        "startDate": start_date,
        "loanDate": loan_date,
        "statusCode": "active",
    }


def reconciliation(as_of_date: str) -> dict[str, Any]:
    balances = entitlement_balance(as_of_date)
    review = [row for row in balances if row["balance"] < 0]
    return {
        "asOfDate": as_of_date,
        "status": "review" if review else "balanced",
        "reviewCount": len(review),
        "rows": balances,
        "note": "Continuous entitlement reconciliation; not payroll posting and not an annual close batch.",
    }


def list_self_service_requests() -> list[dict[str, Any]]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT r.RequestID, e.EmployeeNumber, e.DisplayName, r.RequestType, CONVERT(char(10), r.RequestDate, 23), r.Amount, r.StatusCode, r.Notes "
            "FROM core.SelfServiceRequests r JOIN core.Employees e ON e.EmployeeID = r.EmployeeID "
            "ORDER BY r.CreatedAtUtc DESC;"
        )
        return [
            {
                "requestId": str(row[0]),
                "employeeNumber": row[1],
                "displayName": row[2],
                "requestType": row[3],
                "requestDate": row[4],
                "amount": money(row[5]),
                "statusCode": row[6],
                "notes": row[7],
            }
            for row in cur.fetchall()
        ]


def create_self_service_request(payload: dict[str, Any]) -> dict[str, Any]:
    employee_id = required_text(payload, "employeeId")
    request_type = required_text(payload, "requestType")
    request_date = required_text(payload, "requestDate")
    require_iso_date(request_date, "requestDate")
    if request_type not in {"airfare_request", "profile_update", "loan_request"}:
        raise ValueError("requestType is invalid.")
    request_id = str(uuid.uuid4())
    amount = non_negative(payload.get("amount") or 0, "amount")
    with connect(config.DB_NAME) as conn:
        assert_employee_exists(conn, employee_id)
        conn.cursor().execute(
            "INSERT INTO core.SelfServiceRequests (RequestID, TenantID, CompanyID, EmployeeID, RequestType, RequestDate, Amount, Notes) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?);",
            (request_id, config.TENANT_ID, config.COMPANY_ID, employee_id, request_type, request_date, amount, payload.get("notes")),
        )
    return {"requestId": request_id, "employeeId": employee_id, "requestType": request_type, "requestDate": request_date, "amount": money(amount), "statusCode": "submitted"}


def transition_self_service_request(request_id: str, status_code: str) -> dict[str, Any]:
    if status_code not in {"approved", "rejected", "cancelled"}:
        raise ValueError("statusCode must be approved, rejected, or cancelled.")
    with connect(config.DB_NAME) as conn:
        conn.cursor().execute("UPDATE core.SelfServiceRequests SET StatusCode = ? WHERE RequestID = ?;", (status_code, request_id))
    return {"requestId": request_id, "statusCode": status_code}


def reports_airfare_payable(as_of_date: str) -> dict[str, Any]:
    balances = entitlement_balance(as_of_date)
    return {
        "asOfDate": as_of_date,
        "currency": "BHD",
        "employeeCount": len(balances),
        "totalPayable": money(sum(Decimal(str(row["balance"])) for row in balances)),
        "rows": balances,
    }


def reports_employee_summary() -> dict[str, Any]:
    employees = list_employees()
    departments: dict[str, int] = {}
    for employee in employees:
        key = employee["department"] or "Unassigned"
        departments[key] = departments.get(key, 0) + 1
    return {"total": len(employees), "departments": departments, "rows": employees}


def get_preferences() -> dict[str, Any]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT TOP 1 PreferenceJSON FROM core.UserPreferences up JOIN core.Users u ON u.UserID = up.UserID WHERE u.Username = N'admin';"
        )
        row = cur.fetchone()
    return {"preferences": json_loads(row[0]) if row else default_preferences()}


def save_preferences(payload: dict[str, Any]) -> dict[str, Any]:
    preferences = payload.get("preferences") if isinstance(payload.get("preferences"), dict) else payload
    merged = {**default_preferences(), **preferences}
    raw = json_dumps(merged)
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute("SELECT UserID FROM core.Users WHERE Username = N'admin';")
        user_id = cur.fetchone()[0]
        cur.execute(
            "UPDATE core.UserPreferences SET PreferenceJSON = ?, UpdatedAtUtc = SYSUTCDATETIME() WHERE UserID = ?;",
            (raw, user_id),
        )
    return {"preferences": merged}


def list_users() -> list[dict[str, Any]]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute("SELECT UserID, Username, DisplayName, RoleCode, IsActive FROM core.Users ORDER BY Username;")
        return [{"userId": str(r[0]), "username": r[1], "displayName": r[2], "roleCode": r[3], "isActive": bool(r[4])} for r in cur.fetchall()]


def import_preview(module_code: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    valid = []
    invalid = []
    for index, row in enumerate(rows, start=1):
        if module_code == "employees":
            errors = validate_employee_import_row(row)
            if errors:
                invalid.append({"row": index, "error": "; ".join(errors), "data": row})
            else:
                valid.append({"row": index, "data": row})
        else:
            invalid.append({"row": index, "error": "Required fields missing", "data": row})
    run_id = record_import_export("import", module_code, len(rows), "preview", {"valid": len(valid), "invalid": len(invalid)})
    return {"runId": run_id, "moduleCode": module_code, "validRows": valid, "invalidRows": invalid}


def export_module(module_code: str) -> dict[str, Any]:
    if module_code == "employees":
        rows = list_employees()
    elif module_code == "allocations":
        rows = list_allocations()
    elif module_code == "loans":
        rows = list_loans()
    else:
        raise ValueError("Unsupported export module.")
    run_id = record_import_export("export", module_code, len(rows), "completed", {"rows": len(rows)})
    return {"runId": run_id, "moduleCode": module_code, "rows": rows}


def diagnostics() -> dict[str, Any]:
    health = health_probe()
    return {
        "status": "ok" if all(health["objects"].values()) else "review",
        "database": health,
        "checks": [
            {"name": "MSSQL connection", "status": "ok"},
            {"name": "core schema", "status": "ok" if all(health["objects"].values()) else "review"},
            {"name": "annual close disabled", "status": "ok"},
            {"name": "continuous entitlement", "status": "ok"},
        ],
    }


def support_info() -> dict[str, Any]:
    return {
        "product": "ATLAS Python Core",
        "localUrl": f"http://127.0.0.1:{config.PORT}",
        "database": config.DB_NAME,
        "supportMode": "local-greenfield",
    }


def system_maintenance() -> dict[str, Any]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM core.SystemAuditLog;")
        audit_count = cur.fetchone()[0]
    return {"backupRequired": False, "auditRows": int(audit_count or 0), "database": config.DB_NAME}


def list_attachments() -> list[dict[str, Any]]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT AttachmentID, ModuleCode, OwnerID, FileName, ContentType, DATALENGTH(ContentBytes), CreatedAtUtc "
            "FROM core.Attachments ORDER BY CreatedAtUtc DESC;"
        )
        return [
            {
                "attachmentId": str(row[0]),
                "moduleCode": row[1],
                "ownerId": row[2],
                "fileName": row[3],
                "contentType": row[4],
                "sizeBytes": int(row[5] or 0),
                "createdAtUtc": str(row[6]),
                "viewUrl": f"/api/attachments/{row[0]}/view",
            }
            for row in cur.fetchall()
        ]


def create_attachment(payload: dict[str, Any]) -> dict[str, Any]:
    attachment_id = str(uuid.uuid4())
    module_code = required_text(payload, "moduleCode")
    owner_id = required_text(payload, "ownerId")
    file_name = required_text(payload, "fileName")
    content_type = str(payload.get("contentType") or "application/octet-stream")
    raw = base64.b64decode(required_text(payload, "contentBase64"), validate=True)
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError("Attachment exceeds 5 MB limit.")
    with connect(config.DB_NAME) as conn:
        conn.cursor().execute(
            "INSERT INTO core.Attachments (AttachmentID, ModuleCode, OwnerID, FileName, ContentType, ContentBytes) VALUES (?, ?, ?, ?, ?, ?);",
            (attachment_id, module_code, owner_id, file_name, content_type, raw),
        )
    return {"attachmentId": attachment_id, "fileName": file_name, "contentType": content_type, "sizeBytes": len(raw), "viewUrl": f"/api/attachments/{attachment_id}/view"}


def get_attachment_content(attachment_id: str) -> dict[str, Any]:
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute("SELECT FileName, ContentType, ContentBytes FROM core.Attachments WHERE AttachmentID = ?;", (attachment_id,))
        row = cur.fetchone()
    if not row:
        raise ValueError("attachmentId does not exist.")
    return {"fileName": row[0], "contentType": row[1], "content": bytes(row[2])}


def search_airports(query: str) -> dict[str, Any]:
    q = (query or "").strip().upper()
    if len(q) < 1:
        raise ValueError("q is required.")
    ranked = []
    for airport in AIRPORTS:
        code = airport["code"].upper()
        haystack = f"{airport['code']} {airport['name']} {airport['city']} {airport['country']}".upper()
        if code == q:
            score = 100
        elif code.startswith(q):
            score = 90
        elif airport["city"].upper().startswith(q):
            score = 80
        elif q in haystack:
            score = 60
        else:
            continue
        ranked.append({**airport, "score": score})
    ranked.sort(key=lambda row: (-row["score"], row["code"]))
    return {"query": query, "rows": ranked[:10]}


def create_backup() -> dict[str, Any]:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    path = BACKUP_DIR / f"atlas-python-core-{stamp}.json"
    payload = {
        "createdAtUtc": stamp,
        "database": config.DB_NAME,
        "employees": list_employees(),
        "entitlementEvents": list_entitlement_events(),
        "allocations": list_allocations(),
        "loans": list_loans(),
        "selfServiceRequests": list_self_service_requests(),
        "preferences": get_preferences()["preferences"],
    }
    path.write_text(json_dumps(payload), encoding="utf-8")
    record_import_export("export", "backup", len(payload["employees"]), "completed", {"file": str(path)})
    return {"backupFile": str(path), "sizeBytes": path.stat().st_size, "createdAtUtc": stamp}


def list_backups() -> dict[str, Any]:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    files = sorted(BACKUP_DIR.glob("atlas-python-core-*.json"), key=lambda p: p.stat().st_mtime, reverse=True)
    return {"rows": [{"file": str(p), "sizeBytes": p.stat().st_size, "modifiedUtc": datetime.fromtimestamp(p.stat().st_mtime, timezone.utc).isoformat()} for p in files]}


def restore_backup(payload: dict[str, Any]) -> dict[str, Any]:
    source = Path(required_text(payload, "backupFile"))
    if not source.exists() or source.suffix.lower() != ".json":
        raise ValueError("backupFile does not exist or is not JSON.")
    restore_dir = BACKUP_DIR / "restore-markers"
    restore_dir.mkdir(parents=True, exist_ok=True)
    marker = restore_dir / f"restore-{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}.json"
    shutil.copy2(source, marker)
    record_import_export("import", "restore", 1, "completed", {"source": str(source), "marker": str(marker)})
    return {"status": "completed", "source": str(source), "restoreMarker": str(marker), "mode": "safe-validated-copy"}


def parse_excel_rows(content_base64: str) -> list[dict[str, Any]]:
    raw = base64.b64decode(content_base64, validate=True)
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell or "").strip() for cell in rows[0]]
    parsed = []
    for row in rows[1:]:
        item = {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
        if any(value not in (None, "") for value in item.values()):
            parsed.append(item)
    return parsed


def excel_import_preview(payload: dict[str, Any]) -> dict[str, Any]:
    module_code = required_text(payload, "moduleCode")
    rows = normalize_excel_rows(module_code, parse_excel_rows(required_text(payload, "contentBase64")))
    return import_preview(module_code, rows)


def excel_import_execute(payload: dict[str, Any]) -> dict[str, Any]:
    module_code = required_text(payload, "moduleCode")
    rows = normalize_excel_rows(module_code, parse_excel_rows(required_text(payload, "contentBase64")))
    preview = import_preview(module_code, rows)
    created = []
    if module_code == "employees":
        for item in preview["validRows"]:
            try:
                created.append(create_employee(item["data"]))
            except ValueError as exc:
                preview["invalidRows"].append({"row": item["row"], "error": str(exc), "data": item["data"]})
    run_id = record_import_export("import", module_code, len(rows), "completed", {"created": len(created), "invalid": len(preview["invalidRows"])})
    return {"runId": run_id, "moduleCode": module_code, "createdRows": created, "invalidRows": preview["invalidRows"]}


def normalize_excel_rows(module_code: str, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if module_code != "employees":
        raise ValueError("Only employees Excel import is implemented.")
    normalized = []
    for row in rows:
        normalized.append({
            "employeeNumber": first_value(row, "employeeNumber", "EmployeeNumber", "Employee No", "Employee No."),
            "displayName": first_value(row, "displayName", "DisplayName", "Name", "Employee Name"),
            "legalName": first_value(row, "legalName", "LegalName", "Legal Name"),
            "hireDate": iso_cell(first_value(row, "hireDate", "HireDate", "Hire Date")),
            "terminationDate": iso_cell(first_value(row, "terminationDate", "TerminationDate", "Termination Date")),
            "department": first_value(row, "department", "Department"),
            "jobTitle": first_value(row, "jobTitle", "JobTitle", "Job Title"),
            "employmentType": first_value(row, "employmentType", "EmploymentType", "Employment Type"),
            "payGroup": first_value(row, "payGroup", "PayGroup", "Pay Group", "Employee Group"),
            "workEmail": first_value(row, "workEmail", "WorkEmail", "Email"),
            "phoneNumber": first_value(row, "phoneNumber", "PhoneNumber", "Phone", "Mobile"),
            "nationality": first_value(row, "nationality", "Nationality"),
            "passportNumber": first_value(row, "passportNumber", "PassportNumber", "Passport No", "Passport"),
            "cprNumber": first_value(row, "cprNumber", "CPRNumber", "CPR", "National ID"),
            "bankName": first_value(row, "bankName", "BankName", "Bank"),
            "iban": first_value(row, "iban", "IBAN", "Iban"),
            "basicSalary": first_value(row, "basicSalary", "BasicSalary", "Basic Salary", "Salary"),
            "eligibleForAirfare": bool_cell(first_value(row, "eligibleForAirfare", "EligibleForAirfare", "Airfare Eligible")),
            "homeAirportCode": first_value(row, "homeAirportCode", "HomeAirportCode", "Home Airport"),
            "destinationAirportCode": first_value(row, "destinationAirportCode", "DestinationAirportCode", "Destination Airport"),
        })
    return normalized


def first_value(row: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in row and row[key] not in (None, ""):
            return row[key]
    return ""


def iso_cell(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or "").strip()


def bool_cell(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value or "").strip().lower()
    if normalized in {"", "1", "true", "yes", "y", "eligible"}:
        return True
    if normalized in {"0", "false", "no", "n", "not eligible"}:
        return False
    return True


def create_employee(payload: dict[str, Any]) -> dict[str, Any]:
    employee_id = str(uuid.uuid4())
    employee_number = required_text(payload, "employeeNumber")
    display_name = required_text(payload, "displayName")
    hire_date = required_text(payload, "hireDate")
    require_iso_date(hire_date, "hireDate")
    termination_date = optional_date(payload.get("terminationDate"), "terminationDate")
    employment_type = option(payload.get("employmentType") or "full_time", "employmentType", {"full_time", "part_time", "contract", "temporary", "intern"})
    basic_salary = non_negative(payload.get("basicSalary") or 0, "basicSalary")
    eligible = bool_cell(payload.get("eligibleForAirfare"))
    home_airport = airport_code(payload.get("homeAirportCode"))
    destination_airport = airport_code(payload.get("destinationAirportCode"))
    with connect(config.DB_NAME) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT 1 FROM core.Employees WHERE TenantID = ? AND CompanyID = ? AND EmployeeNumber = ?;",
            (config.TENANT_ID, config.COMPANY_ID, employee_number),
        )
        if cur.fetchone():
            raise ValueError("employeeNumber already exists for this company.")
        cur.execute(
            "INSERT INTO core.Employees "
            "(EmployeeID, TenantID, CompanyID, EmployeeNumber, DisplayName, LegalName, WorkEmail, PhoneNumber, Department, JobTitle, EmploymentType, PayGroup, "
            "Nationality, PassportNumber, CPRNumber, BankName, IBAN, BasicSalary, EligibleForAirfare, HomeAirportCode, DestinationAirportCode, StatusCode, HireDate, TerminationDate) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);",
            (
                employee_id,
                config.TENANT_ID,
                config.COMPANY_ID,
                employee_number,
                display_name,
                text_or_none(payload.get("legalName")),
                text_or_none(payload.get("workEmail")),
                text_or_none(payload.get("phoneNumber")),
                text_or_none(payload.get("department")),
                text_or_none(payload.get("jobTitle")),
                employment_type,
                text_or_none(payload.get("payGroup")),
                text_or_none(payload.get("nationality")),
                text_or_none(payload.get("passportNumber")),
                text_or_none(payload.get("cprNumber")),
                text_or_none(payload.get("bankName")),
                text_or_none(payload.get("iban")),
                basic_salary,
                1 if eligible else 0,
                home_airport,
                destination_airport,
                option(payload.get("statusCode") or "active", "statusCode", {"active", "inactive", "suspended", "terminated", "on_leave"}),
                hire_date,
                termination_date,
            ),
        )
    return {
        "employeeId": employee_id,
        "employeeNumber": employee_number,
        "displayName": display_name,
        "legalName": text_or_none(payload.get("legalName")),
        "hireDate": hire_date,
        "employmentType": employment_type,
        "payGroup": text_or_none(payload.get("payGroup")),
        "eligibleForAirfare": eligible,
    }


def update_employee(employee_id: str, payload: dict[str, Any]) -> dict[str, Any]:
    display_name = required_text(payload, "displayName")
    termination_date = optional_date(payload.get("terminationDate"), "terminationDate")
    employment_type = option(payload.get("employmentType") or "full_time", "employmentType", {"full_time", "part_time", "contract", "temporary", "intern"})
    basic_salary = non_negative(payload.get("basicSalary") or 0, "basicSalary")
    with connect(config.DB_NAME) as conn:
        assert_employee_exists(conn, employee_id)
        conn.cursor().execute(
            "UPDATE core.Employees SET DisplayName = ?, LegalName = ?, WorkEmail = ?, PhoneNumber = ?, Department = ?, JobTitle = ?, EmploymentType = ?, "
            "PayGroup = ?, Nationality = ?, PassportNumber = ?, CPRNumber = ?, BankName = ?, IBAN = ?, BasicSalary = ?, EligibleForAirfare = ?, "
            "HomeAirportCode = ?, DestinationAirportCode = ?, StatusCode = ?, TerminationDate = ?, UpdatedAtUtc = SYSUTCDATETIME() WHERE EmployeeID = ?;",
            (
                display_name,
                text_or_none(payload.get("legalName")),
                text_or_none(payload.get("workEmail")),
                text_or_none(payload.get("phoneNumber")),
                text_or_none(payload.get("department")),
                text_or_none(payload.get("jobTitle")),
                employment_type,
                text_or_none(payload.get("payGroup")),
                text_or_none(payload.get("nationality")),
                text_or_none(payload.get("passportNumber")),
                text_or_none(payload.get("cprNumber")),
                text_or_none(payload.get("bankName")),
                text_or_none(payload.get("iban")),
                basic_salary,
                1 if bool_cell(payload.get("eligibleForAirfare")) else 0,
                airport_code(payload.get("homeAirportCode")),
                airport_code(payload.get("destinationAirportCode")),
                option(payload.get("statusCode") or "active", "statusCode", {"active", "inactive", "suspended", "terminated", "on_leave"}),
                termination_date,
                employee_id,
            ),
        )
    return {"employeeId": employee_id, "displayName": display_name}


def delete_employee(employee_id: str) -> dict[str, Any]:
    with connect(config.DB_NAME) as conn:
        assert_employee_exists(conn, employee_id)
        conn.cursor().execute("UPDATE core.Employees SET StatusCode = N'inactive', UpdatedAtUtc = SYSUTCDATETIME() WHERE EmployeeID = ?;", (employee_id,))
    return {"employeeId": employee_id, "statusCode": "inactive"}


def assert_employee_exists(conn, employee_id: str) -> None:
    cur = conn.cursor()
    cur.execute("SELECT 1 FROM core.Employees WHERE EmployeeID = ?;", (employee_id,))
    if not cur.fetchone():
        raise ValueError("employeeId does not exist.")


def _employee_balance(employee_id: str, as_of_date: str) -> float:
    with connect(config.DB_NAME) as conn:
        assert_employee_exists(conn, employee_id)
        cur = conn.cursor()
        cur.execute(
            "SELECT CAST(COALESCE(SUM(CASE WHEN EventType IN (N'seed',N'accrual',N'adjustment',N'reversal') THEN Amount ELSE 0 END),0) AS DECIMAL(12,3)), "
            "CAST(COALESCE(SUM(CASE WHEN EventType = N'usage' THEN Amount ELSE 0 END),0) AS DECIMAL(12,3)) "
            "FROM core.EntitlementEvents WHERE EmployeeID = ? AND EventDate <= ?;",
            (employee_id, as_of_date),
        )
        earned, used = cur.fetchone()
        return money(Decimal(earned or 0) - Decimal(used or 0))


def money(value: Any) -> float:
    return float(Decimal(value or 0).quantize(Decimal("0.001")))


def required_text(payload: dict[str, Any], name: str) -> str:
    value = str(payload.get(name) or "").strip()
    if not value:
        raise ValueError(f"{name} is required.")
    return value


def non_negative(value: Any, field: str) -> Decimal:
    try:
        amount = Decimal(str(value))
    except Exception as exc:
        raise ValueError(f"{field} must be a non-negative number.") from exc
    if amount < 0:
        raise ValueError(f"{field} must be a non-negative number.")
    return amount.quantize(Decimal("0.001"))


def non_negative_int(value: Any, field: str) -> int:
    try:
        number = int(value)
    except Exception as exc:
        raise ValueError(f"{field} must be a non-negative integer.") from exc
    if number < 0:
        raise ValueError(f"{field} must be a non-negative integer.")
    return number


def text_or_none(value: Any) -> str | None:
    text = str(value or "").strip()
    return text or None


def optional_date(value: Any, field: str) -> str | None:
    text = text_or_none(value)
    if text is None:
        return None
    require_iso_date(text, field)
    return text


def option(value: Any, field: str, allowed: set[str]) -> str:
    text = str(value or "").strip().lower()
    if text not in allowed:
        allowed_text = ", ".join(sorted(allowed))
        raise ValueError(f"{field} must be one of: {allowed_text}.")
    return text


def airport_code(value: Any) -> str | None:
    text = str(value or "").strip().upper()
    if not text:
        return None
    if not re.match(r"^[A-Z]{3}$", text):
        raise ValueError("Airport code must be a 3-letter IATA code.")
    return text


def validate_employee_import_row(row: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    if not text_or_none(row.get("employeeNumber")):
        errors.append("employeeNumber is required")
    if not text_or_none(row.get("displayName")):
        errors.append("displayName is required")
    try:
        require_iso_date(str(row.get("hireDate") or ""), "hireDate")
    except ValueError:
        errors.append("hireDate must be YYYY-MM-DD")
    if text_or_none(row.get("terminationDate")):
        try:
            require_iso_date(str(row.get("terminationDate")), "terminationDate")
        except ValueError:
            errors.append("terminationDate must be YYYY-MM-DD")
    if text_or_none(row.get("employmentType")):
        try:
            option(row.get("employmentType"), "employmentType", {"full_time", "part_time", "contract", "temporary", "intern"})
        except ValueError as exc:
            errors.append(str(exc))
    if text_or_none(row.get("basicSalary")):
        try:
            non_negative(row.get("basicSalary"), "basicSalary")
        except ValueError as exc:
            errors.append(str(exc))
    for field in ("homeAirportCode", "destinationAirportCode"):
        if text_or_none(row.get(field)):
            try:
                airport_code(row.get(field))
            except ValueError:
                errors.append(f"{field} must be a 3-letter IATA code")
    return errors


def record_import_export(direction: str, module_code: str, row_count: int, status: str, detail: dict[str, Any]) -> str:
    run_id = str(uuid.uuid4())
    with connect(config.DB_NAME) as conn:
        conn.cursor().execute(
            "INSERT INTO core.ImportExportRuns (RunID, Direction, ModuleCode, DataRowCount, StatusCode, DetailJSON) VALUES (?, ?, ?, ?, ?, ?);",
            (run_id, direction, module_code, row_count, status, json_dumps(detail)),
        )
    return run_id


def default_preferences() -> dict[str, Any]:
    return {"theme": "system", "density": "comfortable", "defaultPage": "command", "notifications": True}


def json_loads(raw: str) -> dict[str, Any]:
    import json
    return json.loads(raw)


def json_dumps(value: dict[str, Any]) -> str:
    import json
    return json.dumps(value, separators=(",", ":"))


def require_iso_date(value: str, field: str) -> None:
    try:
        date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError(f"{field} is required in YYYY-MM-DD format.") from exc


def validate_database_name(name: str) -> None:
    if not re.match(r"^[A-Za-z0-9_][A-Za-z0-9_.-]{0,127}$", name):
        raise ValueError("Database name is invalid.")

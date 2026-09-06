"""Employee Focus-style Excel template alignment tests."""

from __future__ import annotations

from io import BytesIO
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from airfare_management.api.main import DEFAULT_COMPANY_ID
from airfare_management.infrastructure.documents import EMPLOYEE_COLUMNS, EMPLOYEE_REQUIRED_COLUMNS

pytestmark = pytest.mark.integration


def test_employee_template_matches_focus_master_columns(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    response = client.get("/v1/templates/employees.xlsx", headers=admin_headers)
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    headers = tuple(cell.value for cell in sheet[1])
    assert headers == EMPLOYEE_COLUMNS
    for required in EMPLOYEE_REQUIRED_COLUMNS:
        assert required in headers
    for focus_col in (
        "arabic_name",
        "cpr_no",
        "passport_no",
        "visa_no",
        "grade",
        "contract_type",
        "airline_sector",
        "travel_class",
        "custom_airfare_rate",
    ):
        assert focus_col in headers


def test_employee_import_accepts_focus_style_columns(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    code = f"FOC{uuid4().hex[:5].upper()}"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(list(EMPLOYEE_COLUMNS))
    row = {column: "" for column in EMPLOYEE_COLUMNS}
    row.update(
        {
            "code": code,
            "full_name": "Focus Import User",
            "arabic_name": "مستخدم",
            "gender": "male",
            "date_of_birth": "1991-02-03",
            "nationality": "Bahraini",
            "origin_country": "Bahrain",
            "cpr_no": "910203111",
            "email": "focus.import@example.com",
            "company_id": DEFAULT_COMPANY_ID,
            "join_date": "2024-08-01",
            "department": "HR",
            "branch": "HQ",
            "pay_group": "PG1",
            "repair_center": "HQ",
            "designation": "Officer",
            "sub_section": "General",
            "grade": "G4",
            "contract_type": "unlimited",
            "employment_status": "active",
            "monthly_salary": 500,
            "probation_end_date": "2024-11-01",
            "passport_no": "PFOCUS1",
            "passport_expiry": "2031-01-01",
            "visa_no": "VFOCUS1",
            "visa_expiry": "2027-01-01",
            "airline_sector": "BAH-DXB",
            "travel_class": "Economy",
            "last_airticket_date": "2025-01-10",
            "custom_airfare_rate": 140,
            "max_entitlement_cap_rate": 150,
        }
    )
    sheet.append([row[column] for column in EMPLOYEE_COLUMNS])
    buffer = BytesIO()
    workbook.save(buffer)
    payload = buffer.getvalue()

    preview = client.post(
        "/v1/employees/import/preview",
        headers=admin_headers,
        files={
            "file": (
                "employees.xlsx",
                payload,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200, preview.text
    rows = preview.json()["rows"]
    assert rows[0]["severity"] == "READY"
    assert rows[0]["arabic_name"] == "مستخدم"
    assert rows[0]["cpr_no"] == "910203111"
    assert rows[0]["passport_no"] == "PFOCUS1"
    assert rows[0]["action"] == "INSERT"

    commit = client.post(
        "/v1/employees/import/commit",
        headers=admin_headers,
        json={"rows": [{**rows[0], "selected": True}]},
    )
    assert commit.status_code == 200, commit.text
    assert commit.json()["imported"] == 1

    listed = client.get("/v1/employees", headers=admin_headers, params={"search": code})
    assert listed.status_code == 200
    match = next(item for item in listed.json() if item["code"] == code)
    assert match["arabic_name"] == "مستخدم"
    assert match["cpr_no"] == "910203111"
    assert match["grade"] == "G4"
    assert match["travel_class"] == "Economy"

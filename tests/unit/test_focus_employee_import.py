"""Focus Soft Employee Information workbook parsing."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook

from airfare_management.infrastructure.documents import parse_employee_workbook


def _focus_soft_bytes() -> bytes:
    """Build a title-banded Focus Soft Employee Information sheet."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Employee_Information"
    ws["F1"] = "ATLAS ALUMINUM W.L.L. "
    ws["F2"] = "Employee Information"
    ws["F3"] = "NEW"
    ws["F4"] = " As on 03-09-2026"
    headers = [
        "Job BandCode[General]",
        "Code[General]",
        "Date of Joining[General]",
        "Name[General]",
        "National Identifier (Bahrain ID)[General]",
        "Designation [General]",
        "Nationality[Personal Information]",
        "Pay Group[General]",
        "Sub Section[General]",
        "Reporting To[General]",
        "Emp Department[General]",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(6, col, header)
    ws.cell(7, 1, "OB")
    ws.cell(7, 2, "0001")
    ws.cell(7, 3, datetime(2009, 2, 26))
    ws.cell(7, 4, "HASAN ABDULLA MAHMOOD")
    ws.cell(7, 5, "840602499")
    ws.cell(7, 6, "General Manager")
    ws.cell(7, 7, "BAHRAINI")
    ws.cell(7, 8, "New Pay Group")
    ws.cell(7, 9, "MAIN OFFICE")
    # Reporting To left empty on purpose (sparse mid-row)
    ws.cell(7, 11, "Atlas Transfer 1")
    ws.cell(8, 1, "LTB")
    ws.cell(8, 2, "0003")
    ws.cell(8, 3, datetime(2010, 1, 15))
    ws.cell(8, 4, "ABDULGHANI AHMED ALZAIMOOR")
    ws.cell(8, 5, "800101234")
    ws.cell(8, 6, "Manager")
    ws.cell(8, 7, "BAHRAINI")
    ws.cell(8, 8, "New Pay Group")
    ws.cell(8, 9, "MAIN OFFICE")
    ws.cell(8, 10, "HASAN ABDULLA MAHMOOD")
    ws.cell(8, 11, "Atlas Transfer 1")
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_parse_focus_soft_employee_information_export() -> None:
    rows = parse_employee_workbook(_focus_soft_bytes())
    assert len(rows) == 2
    first = rows[0]
    assert first["code"] == "0001"
    assert first["full_name"] == "HASAN ABDULLA MAHMOOD"
    assert first["join_date"] == date(2009, 2, 26)
    assert first["grade"] == "OB"
    assert first["cpr_no"] == "840602499"
    assert first["designation"] == "General Manager"
    assert first["nationality"] == "BAHRAINI"
    assert first["pay_group"] == "New Pay Group"
    assert first["sub_section"] == "MAIN OFFICE"
    assert first["department"] == "Atlas Transfer 1"
    assert first["company_id"] == "11111111-1111-1111-1111-111111111111"
    assert "reporting_officer_id" not in first or first.get("reporting_officer_id") in (None, "")
    assert rows[1]["reporting_officer_id"] == "HASAN ABDULLA MAHMOOD"
    assert rows[1]["_row"] == 8

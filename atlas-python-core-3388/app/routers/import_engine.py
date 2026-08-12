from __future__ import annotations

import csv
import hashlib
import io
import json
import secrets
from datetime import date, datetime, timezone
from datetime import timedelta
from decimal import Decimal
from enum import StrEnum
from typing import Annotated, Any

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, status
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.database import get_db


router = APIRouter(prefix="/api/v1", tags=["Employee Import"] )
SessionDep = Annotated[Session, Depends(get_db)]

class ImportModule(StrEnum):
    EMPLOYEES = "employees"


class ImportPreviewResponse(BaseModel):
    total: int
    valid: int
    failed: int
    errors: list[dict[str, Any]]
    rows: list[dict[str, Any]]


class ImportFieldStatus(BaseModel):
    column: str
    value: str | None
    status: str = Field(pattern="^(VALID|WARNING|ERROR)$")
    reason: str | None = None


class ImportPreviewRecord(BaseModel):
    row: int
    status: str = Field(pattern="^(VALID|WARNING|ERROR)$")
    fields: list[ImportFieldStatus]
    data: dict[str, Any]


class ImportVerifyPreviewResponse(BaseModel):
    previewToken: str
    expiresAtUtc: datetime
    totalRows: int
    validRowsCount: int
    errorRowsCount: int
    duplicateRowsCount: int
    previewData: list[ImportPreviewRecord]
    errorDetails: list[dict[str, Any]]


class ImportCommitRequest(BaseModel):
    previewToken: str = Field(min_length=32, max_length=96)


class ImportCommitResponse(BaseModel):
    previewToken: str
    insertedRows: int
    skippedRows: int
    status: str


@router.post("/imports/{module_code}/preview", response_model=ImportPreviewResponse)
async def preview_import(module_code: ImportModule, file: UploadFile) -> ImportPreviewResponse:
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Import file exceeds 10 MB.")
    rows = parse_csv_bytes(content)
    errors = []
    valid_rows = []
    for index, row in enumerate(rows, start=1):
        row_errors = validate_employee_import_row(row)
        if row_errors:
            errors.append({"row": index, "errors": row_errors, "data": row})
        else:
            valid_rows.append(row)
    return ImportPreviewResponse(total=len(rows), valid=len(valid_rows), failed=len(errors), errors=errors, rows=valid_rows)


@router.post("/import/verify-preview", response_model=ImportVerifyPreviewResponse)
async def verify_import_preview(file: UploadFile, session: SessionDep, module_code: ImportModule = ImportModule.EMPLOYEES) -> ImportVerifyPreviewResponse:
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Import file exceeds 10 MB.")

    rows = parse_import_file(file.filename or "import.csv", content)
    existing_codes = {
        str(value).upper()
        for value in session.execute(text("SELECT EmployeeCode FROM core.Employees WHERE IsDeleted = 0")).scalars().all()
    }
    seen_codes: set[str] = set()
    preview_records: list[ImportPreviewRecord] = []
    error_details: list[dict[str, Any]] = []
    valid_rows: list[dict[str, Any]] = []
    duplicate_rows = 0

    for row_number, row in enumerate(rows, start=1):
        field_statuses, row_errors, is_duplicate = validate_employee_preview_row(row, row_number, existing_codes, seen_codes)
        if is_duplicate:
            duplicate_rows += 1
        row_status = "ERROR" if row_errors else ("WARNING" if any(field.status == "WARNING" for field in field_statuses) else "VALID")
        for detail in row_errors:
            error_details.append(detail)
        if not row_errors:
            valid_rows.append(normalize_employee_import_row(row))
        preview_records.append(ImportPreviewRecord(row=row_number, status=row_status, fields=field_statuses, data=row))

    token = secrets.token_urlsafe(48)
    expires_at = datetime.now(timezone.utc) + timedelta(hours=2)
    payload = {
        "moduleCode": module_code.value,
        "fileName": file.filename or "import.csv",
        "validRows": valid_rows,
        "errors": error_details,
        "contentSha256": hashlib.sha256(content).hexdigest(),
    }
    session.execute(
        text(
            """
            INSERT INTO core.ImportPreviewSessions
                (PreviewToken, ModuleCode, FileName, PayloadJSON, TotalRows, ValidRowsCount, ErrorRowsCount, DuplicateRowsCount, ExpiresAtUtc)
            VALUES
                (:token, :module_code, :file_name, :payload_json, :total, :valid, :errors, :duplicates, :expires_at)
            """
        ),
        {
            "token": token,
            "module_code": module_code.value,
            "file_name": file.filename or "import.csv",
            "payload_json": json.dumps(payload, separators=(",", ":"), default=str),
            "total": len(rows),
            "valid": len(valid_rows),
            "errors": len(error_details),
            "duplicates": duplicate_rows,
            "expires_at": expires_at.replace(tzinfo=None),
        },
    )
    session.commit()

    return ImportVerifyPreviewResponse(
        previewToken=token,
        expiresAtUtc=expires_at,
        totalRows=len(rows),
        validRowsCount=len(valid_rows),
        errorRowsCount=len(error_details),
        duplicateRowsCount=duplicate_rows,
        previewData=preview_records[:50],
        errorDetails=error_details,
    )


@router.post("/import/commit", response_model=ImportCommitResponse)
def commit_import(payload: ImportCommitRequest, session: SessionDep) -> ImportCommitResponse:
    record = session.execute(
        text(
            """
            SELECT PreviewToken, PayloadJSON, IsCommitted, ExpiresAtUtc
            FROM core.ImportPreviewSessions WITH (UPDLOCK, ROWLOCK)
            WHERE PreviewToken = :token
            """
        ),
        {"token": payload.previewToken},
    ).mappings().first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Preview token not found.")
    if bool(record["IsCommitted"]):
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Preview token was already committed.")
    if record["ExpiresAtUtc"] < datetime.utcnow():
        raise HTTPException(status_code=status.HTTP_410_GONE, detail="Preview token expired. Re-run verification.")

    stored_payload = json.loads(record["PayloadJSON"])
    valid_rows = stored_payload.get("validRows", [])
    inserted = 0
    skipped = 0
    batch_id = session.execute(
        text(
            """
            INSERT INTO core.ImportBatches (ModuleCode, FileName, TotalRows, BatchStatus, CreatedBy)
            OUTPUT INSERTED.ImportBatchID
            VALUES (:module_code, :file_name, :total_rows, N'Processing', N'import-preview')
            """
        ),
        {"module_code": stored_payload.get("moduleCode", "employees"), "file_name": stored_payload.get("fileName", "import.csv"), "total_rows": len(valid_rows)},
    ).scalar_one()

    try:
        for row in valid_rows:
            exists = session.execute(text("SELECT 1 FROM core.Employees WHERE EmployeeCode = :code AND IsDeleted = 0"), {"code": row["EmployeeCode"]}).first()
            if exists:
                skipped += 1
                continue
            session.execute(EMPLOYEE_IMPORT_INSERT_SQL, row)
            inserted += 1
        session.execute(
            text(
                """
                UPDATE core.ImportPreviewSessions SET IsCommitted = 1, CommittedAtUtc = SYSUTCDATETIME() WHERE PreviewToken = :token;
                UPDATE core.ImportBatches SET InsertedRows = :inserted, FailedRows = :skipped, BatchStatus = N'Completed' WHERE ImportBatchID = :batch_id;
                """
            ),
            {"token": payload.previewToken, "inserted": inserted, "skipped": skipped, "batch_id": batch_id},
        )
        session.commit()
    except Exception:
        session.rollback()
        raise

    return ImportCommitResponse(previewToken=payload.previewToken, insertedRows=inserted, skippedRows=skipped, status="Committed")


def parse_csv_bytes(content: bytes) -> list[dict[str, str]]:
    text_content = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text_content))
    return [{str(k): str(v or "").strip() for k, v in row.items()} for row in reader]


def parse_import_file(file_name: str, content: bytes) -> list[dict[str, str]]:
    suffix = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else "csv"
    if suffix == "csv":
        return parse_csv_bytes(content)
    if suffix in {"xlsx", "xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(cell or "").strip() for cell in rows[0]]
        parsed: list[dict[str, str]] = []
        for values in rows[1:]:
            parsed.append({headers[index]: "" if value is None else str(value).strip() for index, value in enumerate(values) if index < len(headers)})
        return parsed
    raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Only CSV, XLSX, and XLSM imports are supported.")


def validate_employee_import_row(row: dict[str, str]) -> list[str]:
    required = ["EmployeeCode", "FullName", "FirstName", "LastName", "Gender", "DateOfBirth", "Nationality", "JoiningDate", "Designation", "EmploymentType", "PaymentMode"]
    errors = [f"{field} is required" for field in required if not row.get(field)]
    if row.get("PaymentMode") in {"Bank", "WPS"} and not row.get("IBANAccountNumber"):
        errors.append("IBANAccountNumber is required for Bank/WPS payment.")
    return errors


def validate_employee_preview_row(row: dict[str, str], row_number: int, existing_codes: set[str], seen_codes: set[str]) -> tuple[list[ImportFieldStatus], list[dict[str, Any]], bool]:
    required = ["EmployeeCode", "FullName", "FirstName", "LastName", "Gender", "DateOfBirth", "Nationality", "JoiningDate", "Designation", "EmploymentType", "PaymentMode"]
    allowed = {
        "Gender": {"Male", "Female", "Other", "Undisclosed"},
        "EmploymentType": {"Permanent", "Contract", "Probation", "Temporary", "Intern"},
        "PaymentMode": {"Bank", "Cash", "WPS"},
        "Status": {"Active", "Inactive", "Resigned", "Terminated", "OnLeave", ""},
    }
    date_fields = ["DateOfBirth", "JoiningDate", "ProbationEndDate", "ConfirmationDate", "PassportExpiry", "CivilIDExpiry", "VisaExpiry", "LabourCardExpiry", "ResignationDate", "LastWorkingDay"]
    money_fields = ["BasicSalary", "HousingAllowance", "TransportAllowance", "OtherFixedAllowances"]
    fields: list[ImportFieldStatus] = []
    errors: list[dict[str, Any]] = []

    def add_error(column: str, reason: str) -> None:
        errors.append({"row": row_number, "column": column, "reason": reason})

    employee_code = str(row.get("EmployeeCode") or "").strip().upper()
    duplicate = False
    if employee_code:
        if employee_code in existing_codes or employee_code in seen_codes:
            duplicate = True
            add_error("EmployeeCode", "Duplicate ID")
        seen_codes.add(employee_code)

    for column in sorted(set(required + date_fields + money_fields + list(row.keys()))):
        value = row.get(column)
        status_value = "VALID"
        reason = None
        if column in required and not str(value or "").strip():
            status_value = "ERROR"
            reason = "Required"
            add_error(column, "Required")
        elif column in date_fields and str(value or "").strip():
            try:
                date.fromisoformat(str(value).strip()[:10])
            except ValueError:
                status_value = "ERROR"
                reason = "Invalid Format"
                add_error(column, "Invalid Format")
        elif column in money_fields and str(value or "").strip():
            try:
                if Decimal(str(value).strip()) < 0:
                    raise ValueError
            except Exception:
                status_value = "ERROR"
                reason = "Invalid non-negative amount"
                add_error(column, "Invalid non-negative amount")
        elif column in allowed and str(value or "").strip() not in allowed[column]:
            status_value = "ERROR"
            reason = "Invalid value"
            add_error(column, "Invalid value")
        elif column == "EmployeeCode" and duplicate:
            status_value = "ERROR"
            reason = "Duplicate ID"
        elif column not in EMPLOYEE_IMPORT_COLUMNS:
            status_value = "WARNING"
            reason = "Ignored column"
        fields.append(ImportFieldStatus(column=column, value=value, status=status_value, reason=reason))
    return fields, errors, duplicate


def normalize_employee_import_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized = {column: row.get(column) or None for column in EMPLOYEE_IMPORT_COLUMNS}
    normalized["Status"] = normalized.get("Status") or "Active"
    normalized["BasicSalary"] = str(normalized.get("BasicSalary") or "0.000")
    normalized["HousingAllowance"] = str(normalized.get("HousingAllowance") or "0.000")
    normalized["TransportAllowance"] = str(normalized.get("TransportAllowance") or "0.000")
    normalized["OtherFixedAllowances"] = str(normalized.get("OtherFixedAllowances") or "0.000")
    normalized["RehireEligible"] = 1 if str(normalized.get("RehireEligible") or "true").lower() in {"1", "true", "yes", "y"} else 0
    return normalized


EMPLOYEE_IMPORT_COLUMNS = [
    "EmployeeCode", "PunchMachineID", "FullName", "FirstName", "MiddleName", "LastName", "PassportName", "Gender", "DateOfBirth", "Nationality", "Religion", "MaritalStatus",
    "JoiningDate", "ProbationEndDate", "ConfirmationDate", "Designation", "GradeLevel", "EmploymentType", "Status",
    "PersonalEmail", "WorkEmail", "MobileNumber", "EmergencyContactName", "EmergencyContactPhone", "EmergencyContactRelationship", "LocalAddress", "HomeCountryAddress",
    "PassportNumber", "PassportExpiry", "CivilID", "CivilIDExpiry", "VisaNumber", "VisaType", "VisaExpiry", "LabourCardNumber", "LabourCardExpiry",
    "BasicSalary", "HousingAllowance", "TransportAllowance", "OtherFixedAllowances", "PaymentMode", "BankName", "IBANAccountNumber", "SwiftCode",
    "ResignationDate", "LastWorkingDay", "ReasonForLeaving", "RehireEligible",
]


EMPLOYEE_IMPORT_INSERT_SQL = text(
    """
    INSERT INTO core.Employees (
        EmployeeCode, PunchMachineID, FullName, FirstName, MiddleName, LastName, PassportName, Gender, DateOfBirth, Nationality, Religion, MaritalStatus,
        JoiningDate, ProbationEndDate, ConfirmationDate, Designation, GradeLevel, EmploymentType, Status,
        PersonalEmail, WorkEmail, MobileNumber, EmergencyContactName, EmergencyContactPhone, EmergencyContactRelationship, LocalAddress, HomeCountryAddress,
        PassportNumber, PassportExpiry, CivilID, CivilIDExpiry, VisaNumber, VisaType, VisaExpiry, LabourCardNumber, LabourCardExpiry,
        BasicSalary, HousingAllowance, TransportAllowance, OtherFixedAllowances, PaymentMode, BankName, IBANAccountNumber, SwiftCode,
        ResignationDate, LastWorkingDay, ReasonForLeaving, RehireEligible
    )
    VALUES (
        :EmployeeCode, :PunchMachineID, :FullName, :FirstName, :MiddleName, :LastName, :PassportName, :Gender, :DateOfBirth, :Nationality, :Religion, :MaritalStatus,
        :JoiningDate, :ProbationEndDate, :ConfirmationDate, :Designation, :GradeLevel, :EmploymentType, :Status,
        :PersonalEmail, :WorkEmail, :MobileNumber, :EmergencyContactName, :EmergencyContactPhone, :EmergencyContactRelationship, :LocalAddress, :HomeCountryAddress,
        :PassportNumber, :PassportExpiry, :CivilID, :CivilIDExpiry, :VisaNumber, :VisaType, :VisaExpiry, :LabourCardNumber, :LabourCardExpiry,
        :BasicSalary, :HousingAllowance, :TransportAllowance, :OtherFixedAllowances, :PaymentMode, :BankName, :IBANAccountNumber, :SwiftCode,
        :ResignationDate, :LastWorkingDay, :ReasonForLeaving, :RehireEligible
    )
    """
)


def json_dumps(value: dict[str, Any]) -> str:
    return json.dumps(value, separators=(",", ":"), default=str)


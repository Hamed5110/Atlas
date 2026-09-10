"""Unit tests for native finance GL (ticket + loan journals)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from airfare_management.infrastructure import finance_ledger as gl


def _build_session() -> sessionmaker[Session]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    ddl = [
        """
        CREATE TABLE companies (
          id VARCHAR(36) PRIMARY KEY,
          code VARCHAR(40),
          name VARCHAR(200),
          deleted_at DATETIME
        )
        """,
        """
        CREATE TABLE finance_accounts (
          id VARCHAR(36) PRIMARY KEY,
          company_id VARCHAR(36) NOT NULL,
          code VARCHAR(20) NOT NULL,
          name VARCHAR(200) NOT NULL,
          account_type VARCHAR(20) NOT NULL,
          currency VARCHAR(3) NOT NULL DEFAULT 'BHD',
          is_active BOOLEAN NOT NULL DEFAULT 1,
          created_at DATETIME NOT NULL,
          updated_at DATETIME NOT NULL,
          created_by VARCHAR(36),
          updated_by VARCHAR(36),
          deleted_at DATETIME
        )
        """,
        """
        CREATE TABLE finance_journals (
          id VARCHAR(36) PRIMARY KEY,
          company_id VARCHAR(36) NOT NULL,
          entry_date DATE NOT NULL,
          narration VARCHAR(500) NOT NULL,
          source_type VARCHAR(40) NOT NULL,
          source_id VARCHAR(36),
          employee_id VARCHAR(36),
          currency VARCHAR(3) NOT NULL DEFAULT 'BHD',
          total_debit NUMERIC(19,4) NOT NULL,
          total_credit NUMERIC(19,4) NOT NULL,
          status VARCHAR(20) NOT NULL DEFAULT 'posted',
          created_at DATETIME NOT NULL,
          updated_at DATETIME NOT NULL,
          created_by VARCHAR(36),
          updated_by VARCHAR(36),
          deleted_at DATETIME
        )
        """,
        """
        CREATE TABLE finance_journal_lines (
          id VARCHAR(36) PRIMARY KEY,
          journal_id VARCHAR(36) NOT NULL,
          account_id VARCHAR(36) NOT NULL,
          account_code VARCHAR(20) NOT NULL,
          debit NUMERIC(19,4) NOT NULL DEFAULT 0,
          credit NUMERIC(19,4) NOT NULL DEFAULT 0,
          memo VARCHAR(400) NOT NULL DEFAULT '',
          created_at DATETIME NOT NULL
        )
        """,
    ]
    with engine.begin() as conn:
        for stmt in ddl:
            conn.execute(text(stmt))
        conn.execute(
            text("INSERT INTO companies (id, code, name) VALUES (:id, 'ATLAS', 'Atlas')"),
            {"id": "co-1"},
        )
    return sessionmaker(bind=engine, expire_on_commit=False)


def test_seed_coa_inserts_six_accounts() -> None:
    sessions = _build_session()
    with sessions() as session:
        assert gl.tables_ready(session) is True
        accounts = gl.ensure_default_coa(session, "co-1")
        raw = session.connection().execute(
            text("SELECT code FROM finance_accounts ORDER BY code")
        ).fetchall()
        assert [r[0] for r in raw] == ["1000", "1100", "1200", "2100", "4000", "5000"], raw
        assert len(accounts) == 6, accounts
    sessions.kw["bind"].dispose()


def test_ticket_issue_loan_journal_balances() -> None:
    sessions = _build_session()
    with sessions() as session:
        posted = gl.post_ticket_issue(
            session,
            company_id="co-1",
            employee_id=uuid4(),
            ticket_id="t-1",
            entry_date=date(2026, 9, 7),
            company_payout=Decimal("100.000"),
            loan_principal=Decimal("50.000"),
            employee_payable=Decimal("50.000"),
            ticket_cost=Decimal("150.000"),
            excess_handling="CONVERT_TO_LOAN",
        )
        assert posted is not None
        assert posted["total_debit"] == posted["total_credit"]
        assert posted["total_debit"] == Decimal("150.00")
        tb = gl.trial_balance(session, "co-1")
        assert tb["balanced"] is True
        assert tb["total_debit"] == "150.00"
        report = gl.ledger_report(session, "co-1", account_code="1200")
        assert any(e["account_code"] == "1200" for e in report["entries"])
    sessions.kw["bind"].dispose()


def test_unbalanced_journal_rejected() -> None:
    sessions = _build_session()
    with sessions() as session:
        with pytest.raises(Exception):
            gl.post_journal(
                session,
                company_id="co-1",
                entry_date=date(2026, 9, 7),
                narration="bad",
                source_type="test",
                source_id=None,
                employee_id=None,
                lines=[
                    ("4000", Decimal("10"), Decimal("0"), "x"),
                    ("1000", Decimal("0"), Decimal("5"), "y"),
                ],
            )
    sessions.kw["bind"].dispose()


def test_ledger_excel_and_pdf_exports() -> None:
    """Focus-style Excel (2 sheets) + PDF magic bytes from a sample balanced report."""
    from io import BytesIO

    from openpyxl import load_workbook

    from airfare_management.infrastructure.finance_ledger import (
        build_ledger_export_pdf,
        build_ledger_export_xlsx,
    )

    report = {
        "as_of": "2026-09-07",
        "balanced": True,
        "currency": "BHD",
        "total_debit": "150.00",
        "total_credit": "150.00",
        "accounts": [
            {
                "code": "1000",
                "name": "Cash / Bank",
                "account_type": "ASSET",
                "debit": "0.00",
                "credit": "150.00",
                "net": "-150.00",
            },
            {
                "code": "4000",
                "name": "Airfare Ticket Expense",
                "account_type": "EXPENSE",
                "debit": "150.00",
                "credit": "0.00",
                "net": "150.00",
            },
        ],
        "entries": [
            {
                "entry_date": "2026-09-07",
                "account_code": "4000",
                "source_type": "ticket_issue",
                "source_id": "t-1",
                "narration": "Ticket issue t-1",
                "memo": "Company airfare payout",
                "debit": "150.00",
                "credit": "0.00",
            },
            {
                "entry_date": "2026-09-07",
                "account_code": "1000",
                "source_type": "ticket_issue",
                "source_id": "t-1",
                "narration": "Ticket issue t-1",
                "memo": "Bank / cash settlement",
                "debit": "0.00",
                "credit": "150.00",
            },
        ],
    }
    xlsx = build_ledger_export_xlsx(report)
    assert xlsx[:2] == b"PK"
    book = load_workbook(BytesIO(xlsx))
    assert "Trial Balance" in book.sheetnames
    assert "Journal Lines" in book.sheetnames
    assert book["Journal Lines"].max_row >= 3

    pdf = build_ledger_export_pdf(report, company_name="Atlas Aluminum")
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 500


def test_loan_recovery_clears_receivable() -> None:
    sessions = _build_session()
    with sessions() as session:
        gl.post_loan_disbursement(
            session,
            company_id="co-1",
            employee_id=uuid4(),
            loan_id="l-1",
            entry_date=date(2026, 1, 1),
            principal=Decimal("80.000"),
        )
        gl.post_loan_recovery(
            session,
            company_id="co-1",
            employee_id=uuid4(),
            loan_id="l-1",
            payment_id="p-1",
            entry_date=date(2026, 2, 1),
            amount=Decimal("80.000"),
        )
        tb = gl.trial_balance(session, "co-1")
        loan = next(a for a in tb["accounts"] if a["code"] == "1200")
        assert loan["debit"] == "80.00"
        assert loan["credit"] == "80.00"
        assert loan["net"] == "0.00"
    sessions.kw["bind"].dispose()

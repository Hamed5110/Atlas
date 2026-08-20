from __future__ import annotations

import json
from pathlib import Path

import openpyxl


def main() -> None:
    path = Path(r"C:\Users\Hamed Ali Khan\Documents\ATLAS Airfare 2024 .xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    print("SHEETS", wb.sheetnames)
    for ws in wb.worksheets:
        print("---", ws.title, ws.max_row, ws.max_column)
        rows = []
        for row in ws.iter_rows(
            min_row=1,
            max_row=min(15, ws.max_row),
            min_col=1,
            max_col=min(12, ws.max_column),
            values_only=True,
        ):
            rows.append(row)
        print(json.dumps(rows, default=str))


if __name__ == "__main__":
    main()
